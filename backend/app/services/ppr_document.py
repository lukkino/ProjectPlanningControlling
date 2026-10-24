"""Verbali di review PPR (Planning, Execution, Deployment, Release to
Market) a partire dal master corporate MOD-PPR (templates/ppr_template.xlsx).

Ogni documento contiene, oltre a Task Guideline, Cover e Action Items, il
foglio della propria review e quelli delle review precedenti:
- Cover: sempre quella del template, con tutte le righe dei firmatari; i
  nomi, ruolo per ruolo, e la Revision History vengono dall'ultima versione
  dello stesso documento (i nomi mancanti dall'ultimo documento della fase
  precedente); i firmatari non previsti per questa review (note "only for
  ..." del template) diventano "n.a.";
- foglio della review: nuovo, dal template (per il Planning con la sezione
  "Deliverable Check" compilata dal popup);
- fogli delle review precedenti e Action Items: dall'ultimo documento della
  fase precedente (es. per l'Execution dall'ultimo Planning).

Si usa Excel via COM e non openpyxl: openpyxl scarta il logo
nell'intestazione e i menu a tendina dei fogli di review, e non sa copiare
fogli da un file all'altro. Serve quindi Excel installato dove gira il
backend; le generazioni sono una alla volta (una sola istanza di Excel)."""

import re
import shutil
import tempfile
import threading
from contextlib import contextmanager
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import openpyxl

from app import models
from app.services.document_generator import (
    DEV_MANAGER,
    HEAD_OF_SW_INTEGRATED_SYSTEM,
    PRODUCT_COMPLIANCE_ENGINEER_SENIOR,
    PRODUCT_MANAGER_TTP,
    PROJECT_MANAGER,
    SYSTEM_NAME,
    TEMPLATES_DIR,
    _estimate_row_height,
    _increment_code,
)

PPR_TEMPLATE_FILENAME = "ppr_template.xlsx"


class PprError(ValueError):
    """File caricato non adatto (tipo di documento sbagliato, fogli
    mancanti...): il messaggio e' per l'utente."""


@dataclass(frozen=True)
class PprType:
    sheet: str  # foglio della review
    title_word: str  # titolo della Cover
    filename: str  # nome file e codice documento (senza versione)
    label: str


# Nell'ordine delle fasi: ogni documento include i fogli dei precedenti.
PPR_DOCUMENT_TYPES: dict[str, PprType] = {
    "planning": PprType("Planning Review", "PLANNING REVIEW", "TIH-PLANNING-PTBSYS", "Planning"),
    "execution": PprType("Execution Review", "EXECUTION REVIEW", "TIH-EXECUTION-PTBSYS", "Execution"),
    "deployment": PprType("Deployment Review", "DEPLOYMENT REVIEW", "TIH-DEPLOYMENT-PTBSYS", "Deployment"),
    "release-to-market": PprType(
        "Release to Market Review", "RELEASE TO MARKET REVIEW", "TIH-RELEASE_TO_MARKET-PTBSYS", "Release to Market"
    ),
}
PHASES = list(PPR_DOCUMENT_TYPES)
TASK_GUIDELINE = "Task Guideline"
COVER = "Cover"
ACTION_ITEMS = "Action Items"
# Nascosto, ma da tenere: e' la sorgente dei menu a tendina dei fogli di review.
INPUT_FOR_EXCEL = "Input for Excel"
NOT_APPLICABLE = "n.a."
# Nel pie' di pagina: segnaposto del template (in arancione) o codice di un
# documento PPR, col colore che lo precede; il codice va in nero.
DOCUMENT_CODE_RE = re.compile(r"(&K[0-9A-Fa-f]{6})?(Document code|TIH-[A-Z_]+-PTBSYS\.\d+)")
FOOTER_TEXT_COLOR = "&K000000"

# Firmatari della Cover, riconosciuti dall'inizio dell'etichetta in colonna
# B (minuscolo, spazi compattati): per quali review servono (dalle note
# "only for ..." del template) e il nome proposto se nessun documento
# caricato ne ha uno. Quelli mai applicabili a questa linea di prodotto
# (Pre-Serie, progetti AP, Mechatronics, rilasci di terze parti) sono
# sempre "n.a.". Attenzione all'ordine: "product manager" non deve
# catturare "production ..." o "product development ...", che iniziano
# diversamente.
_ALL = frozenset(PHASES)
_PLANNING_DEPLOYMENT_RTM = frozenset({"planning", "deployment", "release-to-market"})
SIGNATORY_ROLES: list[tuple[str, frozenset[str], str]] = [
    ("project manager", _ALL, PROJECT_MANAGER),
    ("product compliance engineer", _ALL, PRODUCT_COMPLIANCE_ENGINEER_SENIOR),
    ("independent reviewer", _ALL, ""),
    ("product manager", _PLANNING_DEPLOYMENT_RTM, PRODUCT_MANAGER_TTP),
    ("production quality manager", frozenset(), NOT_APPLICABLE),
    ("product development manager", _PLANNING_DEPLOYMENT_RTM, DEV_MANAGER),
    ("head of development", frozenset(), NOT_APPLICABLE),
    ("head of sw", _PLANNING_DEPLOYMENT_RTM, HEAD_OF_SW_INTEGRATED_SYSTEM),
    ("head of mechatronics", frozenset(), NOT_APPLICABLE),
    ("medical affairs", frozenset({"deployment"}), ""),
    ("third party", frozenset(), NOT_APPLICABLE),
]

# Sezione "Deliverable Check" del foglio Planning Review: righe trovate per
# nome (colonna A), perche' nei documenti compilati si aggiungono righe
# sopra. Colonna F = file (codice e versione), H = note.
PPR_DELIVERABLES: list[str] = [
    "User Needs",
    "Product Qualification and Classification",
    "Product Description",
    "Change Order draft",
    "Change Impact Assessment",
    "Performance Evaluation Plan",
    "Essential Requirement Checklist",
    "Design Input",
    "Risk Management Plan",
    "Safety FMEA",
    "Safety Risk Analysis",
    "Security Risk Analysis",
    "Product Lifecycle Risk Analysis",
    "Production Process Risk Analysis",
    "Criticality Matrix",
    "Design FMEA",
    "Any other additional deliverable",
]
# Testi guida del template nelle celle File/Note, da non riproporre come
# valori nel popup.
_PLACEHOLDER_PREFIXES = ("provide ", "report here")


def _norm(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _role_of(label) -> tuple[str, frozenset[str], str] | None:
    text = _norm(label)
    return next((role for role in SIGNATORY_ROLES if text.startswith(role[0])), None)


def _deliverable_of(label) -> int | None:
    first_line = _norm(str(label or "").split("\n")[0])
    return next((i for i, name in enumerate(PPR_DELIVERABLES) if first_line.startswith(name.lower())), None)


# ---------- Lettura dei documenti caricati (openpyxl, sola lettura) ----------


@dataclass
class _CoverInfo:
    title: str
    names: dict[str, str]  # ruolo -> nome
    history: list[tuple[int, str, str]]  # (versione, autore, descrizione)
    last_version: int | None
    last_description: str | None
    # Documento col vecchio template (prima di MOD-PPR): Cover con altri
    # firmatari, si riconosce dalla mancanza della riga Independent Reviewer.
    old_template: bool


def _is_title(value) -> bool:
    """Titolo della Cover (es. "EXECUTION REVIEW\n PROTUBE SYSTEM - ..."; nel
    template "PLANNING/EXECUTION REVIEW/..."): non basta cercare "REVIEW",
    c'e' anche "This document has been reviewed...". """
    return isinstance(value, str) and any(t.title_word in value.upper() for t in PPR_DOCUMENT_TYPES.values())


def _load(content: bytes, what: str):
    try:
        return openpyxl.load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:  # file non xlsx o danneggiato
        raise PprError(f"{what}: il file non è un documento Excel (.xlsx) leggibile") from exc


def _read_cover(wb, what: str) -> _CoverInfo:
    if COVER not in wb.sheetnames:
        raise PprError(f'{what}: manca il foglio "{COVER}"')
    ws = wb[COVER]
    title = next((str(c.value) for row in ws.iter_rows(max_row=10) for c in row if _is_title(c.value)), "")
    names: dict[str, str] = {}
    roles: set[str] = set()
    history: list[tuple[int, str, str]] = []
    in_history = False
    for row in ws.iter_rows(min_col=1, max_col=4):
        a, b, _, d = (c.value for c in row)
        role = _role_of(b)
        if role:
            roles.add(role[0])
        if role and isinstance(d, str) and d.strip():
            names.setdefault(role[0], d.strip())
        if _norm(a) == "author & revision history":
            in_history = True
        elif in_history and isinstance(a, (int, float)):
            history.append((int(a), str(b or "").strip(), str(d or "").strip()))
    last = history[-1] if history else None
    return _CoverInfo(
        title=title,
        names=names,
        history=history,
        last_version=last[0] if last else None,
        last_description=(last[2] or None) if last else None,
        old_template="independent reviewer" not in roles,
    )


def _check_type(info: _CoverInfo, doc_type: str, what: str) -> None:
    expected = PPR_DOCUMENT_TYPES[doc_type]
    title = info.title.upper()
    other = next((t for t in PPR_DOCUMENT_TYPES.values() if t.title_word in title), None)
    if other is None or other.title_word != expected.title_word:
        found = f"sembra un {other.label}" if other else "il titolo della Cover non indica il tipo di review"
        raise PprError(f"{what}: serve un documento {expected.label}, ma {found}")


def _read_deliverables(wb) -> list[dict] | None:
    if PPR_DOCUMENT_TYPES["planning"].sheet not in wb.sheetnames:
        return None
    found: dict[int, dict] = {}
    for a, *_rest in wb[PPR_DOCUMENT_TYPES["planning"].sheet].iter_rows(min_col=1, max_col=8, values_only=True):
        index = _deliverable_of(a)
        if index is None or index in found:
            continue
        file_value, note_value = _rest[4], _rest[6]
        file_text = "" if _norm(file_value).startswith(_PLACEHOLDER_PREFIXES) else str(file_value or "").strip()
        note_text = "" if _norm(note_value).startswith(_PLACEHOLDER_PREFIXES) else str(note_value or "").strip()
        found[index] = {
            "row": index,
            "name": PPR_DELIVERABLES[index],
            "included": file_text.upper() != "N/A",
            "filename": "" if file_text.upper() == "N/A" else file_text,
            "notes": note_text,
        }
    return [found.get(i, _default_deliverable(i)) for i in range(len(PPR_DELIVERABLES))]


def _default_deliverable(index: int) -> dict:
    return {"row": index, "name": PPR_DELIVERABLES[index], "included": True, "filename": "", "notes": ""}


def previous_phase_of(doc_type: str) -> str | None:
    index = PHASES.index(doc_type)
    return PHASES[index - 1] if index > 0 else None


def get_ppr_defaults(project: models.Project, doc_type: str) -> dict:
    """Valori proposti per il popup prima di caricare i documenti: versione
    1 (primo documento di questo tipo) e il testo di revisione nel formato
    della Revision History ("ProTube Increment PTBSYS-03-003")."""
    return {
        "next_version": 1,
        "last_revision_text": f"ProTube Increment {_increment_code(project)}",
        "deliverables": [_default_deliverable(i) for i in range(len(PPR_DELIVERABLES))],
    }


def inspect_ppr_file(doc_type: str, role: str, content: bytes) -> dict:
    """Controlla un documento caricato nel popup e ne estrae cio' che serve
    a proporre i valori. role: "previous" = ultima versione dello stesso
    documento (ultima versione e descrizione della Revision History, e per
    il Planning i Deliverable Check); "previous_phase" = ultimo documento
    della fase precedente (deve contenerne i fogli e gli Action Items)."""
    expected_type = doc_type if role == "previous" else previous_phase_of(doc_type)
    if expected_type is None:
        raise PprError("Il Planning non ha una fase precedente")
    what = f"Ultimo {PPR_DOCUMENT_TYPES[expected_type].label}"
    wb = _load(content, what)
    info = _read_cover(wb, what)
    _check_type(info, expected_type, what)
    if role == "previous_phase":
        required = [PPR_DOCUMENT_TYPES[p].sheet for p in PHASES[: PHASES.index(doc_type)]] + [ACTION_ITEMS]
        missing = [s for s in required if s not in wb.sheetnames]
        if missing:
            raise PprError(f"{what}: mancano i fogli {', '.join(missing)}")
    return {
        "title": info.title.replace("\n", " ").strip(),
        "last_version": info.last_version,
        "last_description": info.last_description,
        "deliverables": _read_deliverables(wb) if role == "previous" and doc_type == "planning" else None,
        "old_template": info.old_template,
    }


# ---------- Elenco dei documenti attesi, con l'owner (file .txt) ----------

# Documento del "Deliverable Check" (inizio del nome in colonna A, minuscolo)
# -> Task ID (colonna B) delle attivita' del Task Guideline che lo producono,
# da cui si legge l'owner (colonna D). Scelte a mano leggendo gli output
# delle attivita' (colonna F): a cercare il nome nel testo, documenti come
# "Final Test Report" compaiono in piu' attivita' e altri in nessuna.
DELIVERABLE_TASKS: dict[str, tuple[int, ...]] = {
    # Execution Review
    "input/output mapping": (56,),
    "detailed hardware specification": (61,),
    "module control software design & development plan": (62,),
    "standard firmware documentation": (62,),
    "software design & development plan": (63,),
    "software architecture": (63,),
    "design specification": (64,),
    "final test report": (73,),  # SYSTEM & INTEGRATION TEST: e' li' che si firma
    "technical issue list": (78,),
    "installation checklist": (74,),
    "software safety classifications": (65,),
    "compatibility matrix": (83,),
    "driver-core compatibility": (83,),
    "software core compatibility": (83,),
    "automation/host lis interface specifications": (63,),
    "hw/mcsw/sw tool validation": (66,),
    "other software output documents": (62, 63),
    # Deployment Review
    "bom approved": (80,),
    "mcsw & sw package": (72,),
    "operations manual": (87,),
    "translated operations manuals": (90,),
    "service manual": (88,),
    "production instructions": (89,),
    "final risk evaluation": (91,),
    "validation test report": (94,),
    "performance claim list": (97,),
    "technical files": (92,),
    "other technical files": (93,),
    "change order": (105,),
    "regression analysis": (106,),
    "release report": (106,),
    "release training agenda": (86,),
    "performance evaluation report": (99,),
    "post market performance follow up plan": (101,),
    "post market surveillance plan": (102,),
    "general safety and performance requirement": (103,),
    "device master record": (98,),
    "analyzer compatibility": (107,),
    "product certificate or reports": (104,),
    # Release to Market Review
    "technical service bulleting": (116,),
    "technical data sheet": (111,),
    "product information": (117,),
}
OWNER_NOT_FOUND = "owner non trovato nel Task Guideline"


def _review_deliverables(ws) -> list[tuple[str, str]]:
    """(nome, requisito tra parentesi quadre) dei documenti della sezione
    "Deliverable Check" di un foglio di review, senza la riga generica
    "any other additional deliverable"."""
    items: list[tuple[str, str]] = []
    in_section = False
    for (a,) in ws.iter_rows(min_col=1, max_col=1, values_only=True):
        text = _norm(a)
        if text == "deliverable check":
            in_section = True
        elif in_section and (text.startswith("deliverable check") or text.startswith("acceptance criteria")):
            break
        elif in_section and text not in ("", "name") and not text.startswith("any other additional deliverable"):
            name, _, requirement = str(a).partition("\n")
            items.append((re.sub(r"\s+", " ", name).strip(), re.sub(r"\s+", " ", requirement).strip()))
    return items


def _task_owners(ws) -> dict[int, str]:
    owners: dict[int, str] = {}
    for _, task_id, _, owner, *_ in ws.iter_rows(min_row=2, max_col=4, values_only=True):
        if isinstance(task_id, (int, float)) and owner:
            # "DEVELOPMENT TEAM\n(MCSW)" -> "DEVELOPMENT TEAM (MCSW)"; piu'
            # owner su righe diverse -> separati da " / ".
            text = re.sub(r"\s*\n\s*(?=\()", " ", str(owner).strip())
            owners[int(task_id)] = " / ".join(p.strip() for p in text.splitlines() if p.strip())
    return owners


def generate_ppr_deliverables_txt(project: models.Project, doc_type: str) -> tuple[bytes, str]:
    """Elenco dei documenti attesi alla review (Deliverable Check del suo
    foglio nel template), ognuno con l'owner dell'attivita' del Task
    Guideline che lo produce. Si genera insieme al documento PPR."""
    config = PPR_DOCUMENT_TYPES[doc_type]
    wb = openpyxl.load_workbook(TEMPLATES_DIR / PPR_TEMPLATE_FILENAME, data_only=True)
    owners = _task_owners(wb[TASK_GUIDELINE])
    # Una riga per documento, "* documento --> owner": incollata in Jira
    # diventa un elenco puntato. Senza il requisito tra parentesi quadre
    # ("[Document signed]"...).
    lines = [
        f"{config.sheet} - Deliverable Check",
        f"{SYSTEM_NAME} - Increment {_increment_code(project)}",
        "",
    ]
    for name, _requirement in _review_deliverables(wb[config.sheet]):
        key = _norm(name)
        task_ids = next((ids for prefix, ids in DELIVERABLE_TASKS.items() if key.startswith(prefix)), ())
        found = list(dict.fromkeys(owners[t] for t in task_ids if t in owners))
        lines.append(f"* {name} --> {' / '.join(found) if found else OWNER_NOT_FOUND}")
    lines.append("")
    filename = f"{config.sheet.replace(' Review', '').replace(' ', '-')}-review-documents.txt"
    return "\r\n".join(lines).encode("utf-8-sig"), filename


# ---------- Generazione (Excel via COM) ----------

_excel_lock = threading.Lock()

XL_OPEN_XML_WORKBOOK = 51
XL_SHEET_VISIBLE = -1
XL_SHEET_HIDDEN = 0
XL_CELL_TYPE_ALL_VALIDATION = -4174
XL_COLOR_AUTOMATIC = -4105
XL_LEFT = -4131
# Arancione dei segnaposto del template (#EA8651): RGB per openpyxl, BGR per Excel.
PLACEHOLDER_RGB = "EA8651"
PLACEHOLDER_COLOR = 0x5186EA
MSO_AUTOMATION_SECURITY_FORCE_DISABLE = 3


@contextmanager
def _excel():
    """Istanza di Excel nascosta e dedicata (non quella eventualmente aperta
    dall'utente), chiusa in ogni caso alla fine."""
    import pythoncom
    import win32com.client

    with _excel_lock:
        pythoncom.CoInitialize()
        excel = None
        try:
            try:
                excel = win32com.client.DispatchEx("Excel.Application")
            except Exception as exc:
                raise PprError("Excel non è disponibile sul PC del backend: serve per generare i documenti PPR") from exc
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.ScreenUpdating = False
            excel.EnableEvents = False
            excel.AutomationSecurity = MSO_AUTOMATION_SECURITY_FORCE_DISABLE
            yield excel
        finally:
            if excel is not None:
                for wb in list(excel.Workbooks):
                    wb.Close(SaveChanges=False)
                excel.Quit()
            pythoncom.CoUninitialize()


def _copy_validations(source_ws, target_ws) -> None:
    """Copiando un foglio in un altro file Excel scarta i menu a tendina che
    puntano a un altro foglio (qui "Input for Excel"): li ricrea, cella per
    cella, con la stessa formula, che ora punta al foglio del nuovo file."""
    import pywintypes

    try:
        cells = source_ws.Cells.SpecialCells(XL_CELL_TYPE_ALL_VALIDATION)
    except pywintypes.com_error:  # nessuna convalida
        return
    for area in cells.Areas:
        for cell in area.Cells:
            v = cell.Validation
            try:
                formula1 = v.Formula1
            except pywintypes.com_error:
                continue
            try:
                formula2 = v.Formula2
            except pywintypes.com_error:
                formula2 = None
            target = target_ws.Range(cell.Address).Validation
            target.Delete()
            args = {"Type": v.Type, "AlertStyle": v.AlertStyle, "Operator": v.Operator, "Formula1": formula1}
            if formula2:
                args["Formula2"] = formula2
            target.Add(**args)
            target.IgnoreBlank = v.IgnoreBlank
            target.InCellDropdown = v.InCellDropdown
            target.ShowInput = v.ShowInput
            target.ShowError = v.ShowError
            target.InputTitle = v.InputTitle
            target.InputMessage = v.InputMessage
            target.ErrorTitle = v.ErrorTitle
            target.ErrorMessage = v.ErrorMessage


def _replace_sheet(base, source, name: str) -> None:
    source_ws = source.Worksheets(name)
    source_ws.Visible = XL_SHEET_VISIBLE  # un foglio nascosto non si copia
    old = base.Worksheets(name)
    source_ws.Copy(Before=old)
    copied = base.Worksheets(old.Index - 1)
    old.Delete()
    copied.Name = name
    copied.Visible = XL_SHEET_VISIBLE
    _copy_validations(source_ws, copied)


def _write(cell, value) -> None:
    """Valore in una cella (o nell'area unita che la contiene), nel colore
    normale del testo: nel template i segnaposto sono arancioni."""
    cell.MergeArea.ClearContents()
    cell.Value = value
    cell.MergeArea.Font.ColorIndex = XL_COLOR_AUTOMATIC


def _used_rows(ws) -> int:
    used = ws.UsedRange
    return used.Row + used.Rows.Count - 1


def _signatory_name(key: str, previous: dict[str, str], previous_phase: dict[str, str], default: str) -> str:
    """Nome per un ruolo previsto da questa review: quello dell'ultima
    versione dello stesso documento (anche "n.a.", se li' era stato messo
    a mano), altrimenti quello dell'ultimo documento della fase precedente
    (dove "n.a." puo' voler dire solo che in quella review il ruolo non
    c'era), altrimenti il nome proposto."""
    if previous.get(key):
        return NOT_APPLICABLE if previous[key].lower() == NOT_APPLICABLE else previous[key]
    if previous_phase.get(key, "").lower() not in ("", NOT_APPLICABLE):
        return previous_phase[key]
    return default


def _fill_cover(
    cover,
    doc_type: str,
    increment: str,
    previous: _CoverInfo | None,
    previous_phase: _CoverInfo | None,
    sheets: list[str],
    version: int,
    revision_text: str,
) -> None:
    """Cover del template (titolo, tutte le righe dei firmatari previste da
    MOD-PPR, nota sui fogli) con i nomi e la Revision History dei documenti
    caricati. Non si copia la Cover caricata: quelle dei documenti piu'
    vecchi hanno un'altra struttura (es. senza Independent Reviewer)."""
    # Ultima versione col vecchio template (fase di transizione): valgono
    # solo versione e Revision History, i firmatari erano altri.
    previous_names = previous.names if previous and not previous.old_template else {}
    phase_names = previous_phase.names if previous_phase else {}
    config = PPR_DOCUMENT_TYPES[doc_type]
    last_row = _used_rows(cover)
    history_row = None
    for row in range(1, last_row + 1):
        a = cover.Cells(row, 1).Value
        b = cover.Cells(row, 2).Value
        if row <= 10:
            for col in range(1, 13):
                value = cover.Cells(row, col).Value
                if _is_title(value):
                    _write(cover.Cells(row, col), f"{config.title_word}\n {SYSTEM_NAME.upper()} - INCREMENT {increment}")
        role = _role_of(b)
        if role is not None:
            key, phases, default = role
            name = _signatory_name(key, previous_names, phase_names, default) if doc_type in phases else NOT_APPLICABLE
            _write(cover.Cells(row, 4), name)
        if isinstance(a, str) and _norm(a).startswith("note: the complete document includes"):
            _write(cover.Cells(row, 1), f"Note: the complete document includes the following worksheets: {', '.join(sheets)}.")
        if _norm(a) == "author & revision history":
            history_row = row
    if history_row is None:
        raise PprError('Cover: non trovo la tabella "AUTHOR & REVISION HISTORY"')

    # Prima riga dopo l'intestazione (Version/Author/Change Description): nel
    # template e' un segnaposto; le versioni precedenti e la nuova vanno da
    # li' in giu', ognuna su una riga col formato del segnaposto.
    # Tutto in blocco: cella per cella, con una ventina di versioni, Excel
    # impiega decine di secondi.
    first = history_row + 2
    base_height = cover.Rows(first).RowHeight
    entries = [*(previous.history if previous else []), (version, PROJECT_MANAGER, revision_text)]
    last = first + len(entries) - 1
    if last > first:
        cover.Rows(f"{first + 1}:{last}").Insert()
        cover.Rows(first).Copy(cover.Rows(f"{first + 1}:{last}"))
    rows = cover.Range(f"A{first}:L{last}")
    rows.ClearContents()
    rows.Font.ColorIndex = XL_COLOR_AUTOMATIC
    # Il segnaposto del template e' centrato e su una riga: le descrizioni
    # lunghe verrebbero tagliate.
    descriptions = cover.Range(f"D{first}:L{last}")
    descriptions.HorizontalAlignment = XL_LEFT
    descriptions.WrapText = True
    for col, values in (("A", [e[0] for e in entries]), ("B", [e[1] for e in entries]), ("D", [e[2] for e in entries])):
        cover.Range(f"{col}{first}:{col}{last}").Value = [[v] for v in values]
    for i, (_, _, description) in enumerate(entries):
        height = max(base_height, _estimate_row_height(description))
        if height != base_height:
            cover.Rows(first + i).RowHeight = height


def _fill_attendees(sheet, attendees: list[tuple[str, str]]) -> None:
    """Sezione "Meeting Attendees and Roles" del foglio nuovo: i partecipanti
    del meeting precedente al posto delle righe del template, togliendo i
    segnaposto arancioni. Resta solo quello dell'Independent Reviewer (sempre
    previsto) se nel meeting precedente non c'era."""
    header = next(
        (r for r in range(1, 30) if _norm(sheet.Cells(r, 1).Value) == "name & surname"), None
    )
    if header is None:
        return
    end = header + 1
    while end < header + 40 and sheet.Cells(end, 1).MergeArea.Columns.Count < 10:  # titolo della sezione dopo
        end += 1
    first = header + 1

    reviewer_placeholder = None
    for row in range(first, end):
        if "independent reviewer" in _norm(sheet.Cells(row, 1).Value):
            reviewer_placeholder = (sheet.Cells(row, 1).Value, sheet.Cells(row, 6).Value)
    rows: list[tuple[str, str, bool]] = [(name, role, False) for name, role in attendees]
    if reviewer_placeholder and not any("independent reviewer" in _norm(role) for _, role in attendees):
        rows.append((*reviewer_placeholder, True))

    slots = end - first
    if len(rows) > slots:
        extra = f"{end}:{end + len(rows) - slots - 1}"
        sheet.Rows(extra).Insert()
        sheet.Rows(first).Copy(sheet.Rows(extra))
        end += len(rows) - slots
    area = sheet.Range(f"A{first}:J{end - 1}")
    area.ClearContents()
    area.Font.ColorIndex = XL_COLOR_AUTOMATIC
    for i, (name, role, placeholder) in enumerate(rows):
        sheet.Cells(first + i, 1).Value = name
        sheet.Cells(first + i, 6).Value = role
        if placeholder:
            sheet.Cells(first + i, 1).MergeArea.Font.Color = PLACEHOLDER_COLOR


def _read_attendees(wb, sheet_name: str) -> list[tuple[str, str]] | None:
    """Partecipanti (nome, ruolo) di un meeting, saltando i segnaposto
    arancioni del template. None se il foglio o la sezione non ci sono."""
    if sheet_name not in wb.sheetnames:
        return None
    ws = wb[sheet_name]
    header = next(
        (c.row for (c,) in ws.iter_rows(max_row=30, max_col=1) if _norm(c.value) == "name & surname"), None
    )
    if header is None:
        return None
    section_titles = {m.min_row for m in ws.merged_cells.ranges if m.min_col == 1 and m.max_col >= 10}
    attendees: list[tuple[str, str]] = []
    row = header + 1
    while row not in section_titles and row < header + 40:
        name_cell, role = ws.cell(row, 1), ws.cell(row, 6).value
        color = name_cell.font.color
        placeholder = color is not None and color.type == "rgb" and str(color.rgb).upper().endswith(PLACEHOLDER_RGB)
        name = str(name_cell.value or "").strip()
        if name and not placeholder and not _norm(name).startswith("report the name"):
            attendees.append((name, str(role or "").strip()))
        row += 1
    return attendees


def _fill_deliverables(sheet, deliverables: list[dict]) -> None:
    by_index = {d.get("row"): d for d in deliverables}
    for row in range(1, _used_rows(sheet) + 1):
        index = _deliverable_of(sheet.Cells(row, 1).Value)
        item = by_index.get(index) if index is not None else None
        if item is None:
            continue
        included = item.get("included", True)
        filename = (item.get("filename") or "").strip()
        _write(sheet.Cells(row, 6), filename if included else "N/A")
        _write(sheet.Cells(row, 8), (item.get("notes") or "").strip())
        by_index.pop(index)


def generate_ppr_document(
    project: models.Project,
    doc_type: str,
    version: int,
    revision_text: str,
    deliverables: list[dict] | None,
    previous: bytes | None,
    previous_phase: bytes | None,
) -> tuple[bytes, str]:
    """Documento PPR di tipo doc_type (vedi la docstring del modulo).
    previous: ultima versione dello stesso documento (facoltativa: senza, la
    Revision History parte dalla versione scelta); previous_phase: ultimo documento della fase
    precedente, obbligatorio tranne che per il Planning."""
    config = PPR_DOCUMENT_TYPES[doc_type]
    index = PHASES.index(doc_type)
    earlier_sheets = [PPR_DOCUMENT_TYPES[p].sheet for p in PHASES[:index]]
    if index > 0 and previous_phase is None:
        raise PprError(f"Serve l'ultimo documento {PPR_DOCUMENT_TYPES[PHASES[index - 1]].label}")

    # Controlli, nomi dei firmatari, Revision History e partecipanti al
    # meeting precedente prima di aprire Excel.
    previous_cover = phase_cover = None
    attendees = None
    if previous is not None:
        inspect_ppr_file(doc_type, "previous", previous)
        previous_book = _load(previous, "Ultima versione")
        previous_cover = _read_cover(previous_book, "Ultima versione")
        # Del vecchio template valgono solo versione e Revision History.
        if not previous_cover.old_template:
            attendees = _read_attendees(previous_book, config.sheet)
    if previous_phase is not None:
        inspect_ppr_file(doc_type, "previous_phase", previous_phase)
        phase_book = _load(previous_phase, "Fase precedente")
        phase_cover = _read_cover(phase_book, "Fase precedente")
        if attendees is None:
            attendees = _read_attendees(phase_book, PPR_DOCUMENT_TYPES[PHASES[index - 1]].sheet)

    review_sheets = earlier_sheets + [config.sheet]
    keep = [TASK_GUIDELINE, COVER, *review_sheets, ACTION_ITEMS, INPUT_FOR_EXCEL]
    document_id = f"{config.filename}.{version}"

    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp_dir, _excel() as excel:
        tmp = Path(tmp_dir)
        shutil.copy(TEMPLATES_DIR / PPR_TEMPLATE_FILENAME, tmp / "base.xlsx")
        base = excel.Workbooks.Open(str(tmp / "base.xlsx"), UpdateLinks=0)

        def open_source(content: bytes, name: str):
            (tmp / name).write_bytes(content)
            return excel.Workbooks.Open(str(tmp / name), UpdateLinks=0, ReadOnly=True)

        if previous_phase is not None:
            phase_wb = open_source(previous_phase, "previous_phase.xlsx")
            for name in [*earlier_sheets, ACTION_ITEMS]:
                _replace_sheet(base, phase_wb, name)

        for ws in list(base.Worksheets):
            if ws.Name not in keep:
                ws.Visible = XL_SHEET_VISIBLE  # l'ultimo foglio visibile non si cancella, uno nascosto si'
                ws.Delete()
        base.Worksheets(INPUT_FOR_EXCEL).Visible = XL_SHEET_HIDDEN

        cover = base.Worksheets(COVER)
        _fill_cover(
            cover,
            doc_type,
            _increment_code(project),
            previous_cover,
            phase_cover,
            [COVER, *review_sheets, ACTION_ITEMS],
            version,
            revision_text,
        )
        if attendees is not None:
            _fill_attendees(base.Worksheets(config.sheet), attendees)
        if doc_type == "planning":
            _fill_deliverables(base.Worksheets(config.sheet), deliverables or [])

        # Codice documento nel pie' di pagina a sinistra, sotto "MOD-PPR.1":
        # al posto di "Document code" del template o del codice della
        # versione da cui arriva il foglio, lasciando il resto (formattazione
        # compresa). La Task Guideline ha il suo (MOD-PMC).
        excel.PrintCommunication = False
        for name in [COVER, *review_sheets, ACTION_ITEMS]:
            page_setup = base.Worksheets(name).PageSetup
            code = FOOTER_TEXT_COLOR + document_id
            footer, count = DOCUMENT_CODE_RE.subn(lambda _: code, page_setup.LeftFooter or "")
            page_setup.LeftFooter = footer if count else f"{footer}\n{code}".lstrip("\n")
        excel.PrintCommunication = True

        cover.Activate()
        cover.Range("A1").Select()
        out = tmp / "out.xlsx"
        base.SaveAs(str(out), FileFormat=XL_OPEN_XML_WORKBOOK)
        base.Close(SaveChanges=False)
        content = out.read_bytes()
    return content, config.filename
