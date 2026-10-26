"""Presentazioni pptx degli scenari delle Story, nella cartella Scenarios
dell'increment (Project.scenarios_folder, creata a mano dall'utente).

Un file appartiene a una Story se il nome inizia con la sua chiave Jira
(es. "PTBSYS-1234-Scenarios.pptx"): cosi' vale anche per i file creati o
rinominati a mano, purche' mantengano la chiave in testa.

La presentazione nuova si genera da TEMPLATE_PATH (ricavato da
PTBSYS-5872-Scenarios.pptx): cover + una slide scenario di esempio, che fa
da modello e viene ripetuta per ogni "Make sure that" (acceptance criteria)
della description Jira della Story."""

import copy
import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.slide import Slide
from pptx.util import Pt

from app import models
from app.services.document_generator import TEMPLATES_DIR

TEMPLATE_PATH = TEMPLATES_DIR / "scenarios_template.pptx"

# Testi segnaposto delle slide scenario, per le parti di Given/When/Then che
# propose_scenario non riesce a ricavare dal criterio.
GIVEN_PLACEHOLDER = "[precondition]"
WHEN_PLACEHOLDER = "[action]"
THEN_PLACEHOLDER = "[expected result]"
TITLE_PLACEHOLDER = "– [Title]"
# Oltre questa lunghezza la summary sulla cover, a 18pt, scende sopra
# "Last update": la si scrive piu' piccola.
LONG_SUMMARY_CHARS = 150

# Punto elenco dei sotto-punti di un criterio (come in PTBSYS-5872, slide 9).
_SUB_ITEM_PPR = (
    '<a:pPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" marL="742950" lvl="1" indent="-285750">'
    '<a:buFont typeface="Arial" panose="020B0604020202020204" pitchFamily="34" charset="0"/>'
    '<a:buChar char="&#8226;"/></a:pPr>'
)


def _key_prefix_pattern(jira_key: str) -> re.Pattern[str]:
    # La chiave non deve proseguire con un altro carattere alfanumerico:
    # "PTBSYS-12" non deve riconoscere "PTBSYS-123-Scenarios.pptx".
    return re.compile(rf"^{re.escape(jira_key)}(?![A-Za-z0-9])", re.IGNORECASE)


def find_scenario_files(folder: Path, jira_keys: list[str]) -> dict[str, str]:
    """Chiave Jira -> nome del file pptx degli scenari presente in folder.
    Con piu' file per la stessa chiave vince il primo in ordine alfabetico
    (es. "-Scenarios.pptx" prima di "-Scenarios_old.pptx")."""
    names = sorted(
        p.name
        for p in folder.iterdir()
        if p.is_file() and p.suffix.lower() == ".pptx" and not p.name.startswith("~$")  # ~$ = lock di PowerPoint
    )
    found: dict[str, str] = {}
    for key in jira_keys:
        pattern = _key_prefix_pattern(key)
        match = next((n for n in names if pattern.match(n)), None)
        if match:
            found[key] = match
    return found


def scenario_filename(item: models.BacklogItem) -> str:
    return f"{item.jira_key}-Scenarios.pptx"


@dataclass
class AcceptanceCriterion:
    # Testo dopo "Make sure that", senza i marcatori *grassetto* di Jira.
    text: str
    sub_items: list[str] = field(default_factory=list)


_BULLET_LINE = re.compile(r"^(\s*)([-•]|\d+[.)])?\s*(.*)$")
_MAKE_SURE = re.compile(r"^make\s+sure\s+(?:that\b)?\s*(.*)$", re.IGNORECASE)


def parse_acceptance_criteria(description: str | None) -> list[AcceptanceCriterion]:
    """I "Make sure that ..." della description, nell'ordine, ciascuno con i
    suoi sotto-punti (le righe elenco piu' rientrate che lo seguono)."""
    criteria: list[AcceptanceCriterion] = []
    current: AcceptanceCriterion | None = None
    current_indent = 0
    for raw in (description or "").splitlines():
        match = _BULLET_LINE.match(raw.replace("*", ""))
        indent, bullet, text = len(match.group(1)), match.group(2), " ".join(match.group(3).split())
        if not text:
            continue
        make_sure = _MAKE_SURE.match(text)
        if make_sure:
            current = AcceptanceCriterion(make_sure.group(1))
            current_indent = indent
            criteria.append(current)
        elif current is not None and bullet and indent > current_indent:
            current.sub_items.append(text)
        else:
            current = None
    return criteria


@dataclass
class ScenarioDraft:
    # Bozza di Given/When/Then ricavata con regole dal testo del criterio:
    # liste vuote / None = parte non riconosciuta, resta il segnaposto.
    given: list[str] = field(default_factory=list)
    when: str | None = None
    then: list[str] = field(default_factory=list)


# Riferimenti ai rischi in coda al criterio: restano nel riquadro del
# criterio, non servono nello scenario.
_MITIGATES = re.compile(r"\s*\((?:mitigates|expands)\b[^)]*\)", re.IGNORECASE)
# Inizio tipico della conseguenza quando condizione e conseguenza non sono
# separate da una virgola ("if there is no rack ... the system asks ...").
_CONSEQUENCE_START = re.compile(
    r"\s(?=(?:the (?:system|gui|user|application|software)|an? (?:error|warning|message|popup)|it)\b)",
    re.IGNORECASE,
)
# " and " che apre una nuova condizione (non "checked-in and printed tubes").
_AND_CLAUSE = re.compile(r"\s+and\s+(?=(?:the|a|an|no|all|each|there|it)\b)", re.IGNORECASE)


def _split_top_level_comma(text: str) -> tuple[str, str] | None:
    """Divide alla prima virgola fuori dalle parentesi."""
    depth = 0
    for i, ch in enumerate(text):
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth = max(0, depth - 1)
        elif ch == "," and depth == 0:
            return text[:i].strip(), text[i + 1 :].strip()
    return None


def _split_condition(text: str) -> tuple[str, str | None]:
    """"<condizione>, <conseguenza>" -> (condizione, conseguenza); senza
    separatore riconoscibile la conseguenza e' None."""
    parts = _split_top_level_comma(text)
    if parts and parts[1]:
        return parts
    then_split = re.split(r"\s+then\s+", text, maxsplit=1, flags=re.IGNORECASE)
    if len(then_split) == 2:
        return then_split[0].strip(), then_split[1].strip()
    # La condizione e' lunga almeno qualche parola: evita di spezzare su un
    # "it"/"the user" all'inizio.
    for match in _CONSEQUENCE_START.finditer(text):
        if match.start() > 10:
            return text[: match.start()].strip(), text[match.end() :].strip()
    return text, None


def _lower_first(text: str) -> str:
    # Non le sigle: "GUI shows" resta com'e'.
    return text if text[1:2].isupper() else text[:1].lower() + text[1:]


def _sentence(text: str) -> str:
    text = text.strip().rstrip(".").strip()
    return text[:1].upper() + text[1:]


def propose_scenario(criterion: AcceptanceCriterion) -> ScenarioDraft:
    """Given/When/Then proposti dal testo del criterio, senza AI:
    "if|in case of X, Y" -> Given X / Then Y; "after|when X, ..." -> When X; il resto
    del criterio e i suoi sotto-punti vanno nel Then."""
    text = _MITIGATES.sub("", criterion.text).strip()
    draft = ScenarioDraft()

    trigger = re.match(r"^(?:after|when|once)\s+(.*)$", text, re.IGNORECASE)
    if trigger:
        rest = trigger.group(1)
        # "after X, if Y, Z" / "after X if Y Z": X e' il When, il resto si
        # rianalizza come condizione.
        inner_if = re.search(r",?\s+if\s+", rest, re.IGNORECASE)
        if inner_if:
            draft.when = rest[: inner_if.start()].strip(" ,")
            text = "if " + rest[inner_if.end() :]
        else:
            when, consequence = _split_condition(rest)
            if consequence:
                draft.when, text = when, consequence

    condition = re.match(r"^(?:if|in case of)\s+(.*)$", text, re.IGNORECASE)
    if condition:
        given, consequence = _split_condition(condition.group(1))
        # Minuscola: segue "Given"/"And" sulla slide.
        draft.given = [_lower_first(s) for s in (_sentence(g) for g in _AND_CLAUSE.split(given)) if s]
        text = consequence or ""

    if text:
        draft.then.append(_sentence(text))
    draft.then.extend(_sentence(_MITIGATES.sub("", s)) for s in criterion.sub_items)
    if draft.when:
        draft.when = _lower_first(draft.when)
    return draft


def _duplicate_slide(prs: Presentation, source: Slide) -> Slide:
    # La slide modello ha solo la relazione col layout (nessuna immagine):
    # basta copiarne le forme.
    slide = prs.slides.add_slide(source.slide_layout)
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    for element in source.shapes._spTree.iterchildren():
        if element.tag in (qn("p:sp"), qn("p:grpSp"), qn("p:graphicFrame"), qn("p:cxnSp"), qn("p:pic")):
            slide.shapes._spTree.append(copy.deepcopy(element))
    return slide


def _remove_slide(prs: Presentation, slide: Slide) -> None:
    slide_ids = prs.slides._sldIdLst
    for slide_id in slide_ids:
        if prs.slides.get(int(slide_id.get("id"))) is slide:
            prs.part.drop_rel(slide_id.rId)
            slide_ids.remove(slide_id)
            return


def _shape(slide: Slide, name: str):
    return next(s for s in slide.shapes if s.name == name)


def _fill_cover(slide: Slide, item: models.BacklogItem, project: models.Project) -> None:
    title_runs = slide.shapes.title.text_frame.paragraphs[0].runs
    title_runs[1].text = f" for {item.jira_key}"
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        runs = shape.text_frame.paragraphs[0].runs
        text = shape.text_frame.text
        if shape.name == "TextBox 2":
            runs[0].text = item.summary or ""
            if len(runs[0].text) > LONG_SUMMARY_CHARS:
                runs[0].font.size = Pt(14)
        elif text.startswith("Last update"):
            runs[0].text = f"Last update {dt.date.today():%d/%m/%Y}"
        elif text.startswith("Increment"):
            runs[0].text = f"Increment {project.code}"


def _repeat_paragraph(paragraph, texts: list[str]) -> None:
    """Una copia del paragrafo (stessa formattazione, testo in un solo run)
    per ogni testo, al suo posto; nessun testo = paragrafo rimosso."""
    anchor = paragraph._p
    for text in texts:
        copy_p = copy.deepcopy(paragraph._p)
        runs = copy_p.findall(qn("a:r"))
        for extra in runs[1:]:
            copy_p.remove(extra)
        runs[0].find(qn("a:t")).text = text
        anchor.addnext(copy_p)
        anchor = copy_p
    paragraph._p.getparent().remove(paragraph._p)


def _fill_scenario(slide: Slide, item: models.BacklogItem, number: int, criterion: AcceptanceCriterion | None) -> None:
    title_runs = slide.shapes.title.text_frame.paragraphs[0].runs
    title_runs[0].text = item.jira_key
    title_runs[2].text = f"{number} "
    title_runs[3].text = TITLE_PLACEHOLDER

    # Modello: Given, And x N, When, Then, risultato. Si tiene una riga per
    # tipo e si ripetono "And" e risultato quante volte servono.
    draft = propose_scenario(criterion) if criterion else ScenarioDraft()
    body = _shape(slide, "Content Placeholder 2").text_frame
    paragraphs = body.paragraphs
    for extra in paragraphs[2:-3]:
        extra._p.getparent().remove(extra._p)
    given, and_, when, _then, result = body.paragraphs
    given_lines = draft.given or [GIVEN_PLACEHOLDER]
    given.runs[1].text = f" {given_lines[0]}"
    _repeat_paragraph(and_, [f"And {line}" for line in given_lines[1:]])
    when.runs[1].text = f" {draft.when or WHEN_PLACEHOLDER}"
    _repeat_paragraph(result, draft.then or [THEN_PLACEHOLDER])

    # Criterio di accettazione in fondo: "<n> - Make sure that" in grassetto,
    # poi il testo e gli eventuali sotto-punti.
    ac_frame = _shape(slide, "TextBox 3").text_frame
    first = ac_frame.paragraphs[0]
    bold_rpr = copy.deepcopy(first.runs[0]._r.rPr)
    plain_rpr = copy.deepcopy(next(r for r in first.runs if not r.font.bold)._r.rPr)
    for extra in ac_frame.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    for run in first.runs:
        run._r.getparent().remove(run._r)

    def add_run(paragraph, text: str, rpr) -> None:
        run = paragraph.add_run()
        run.text = text
        run._r.insert(0, copy.deepcopy(rpr))

    if criterion is None:
        add_run(first, f"{number} - Make sure that", bold_rpr)
        add_run(first, " [acceptance criterion]", plain_rpr)
        return
    add_run(first, f"{number} - Make sure that", bold_rpr)
    add_run(first, f" {criterion.text}", plain_rpr)
    for sub_item in criterion.sub_items:
        paragraph = ac_frame.add_paragraph()
        paragraph._p.insert(0, etree.fromstring(_SUB_ITEM_PPR))
        add_run(paragraph, sub_item, plain_rpr)


def create_scenario_presentation(folder: Path, item: models.BacklogItem, project: models.Project) -> str:
    """Crea la presentazione degli scenari della Story (una slide per
    criterio di accettazione, almeno una) e ne restituisce il nome del file."""
    prs = Presentation(str(TEMPLATE_PATH))
    cover, model_slide = prs.slides[0], prs.slides[1]
    _fill_cover(cover, item, project)
    criteria: list[AcceptanceCriterion | None] = list(parse_acceptance_criteria(item.description)) or [None]
    for number, criterion in enumerate(criteria, start=1):
        _fill_scenario(_duplicate_slide(prs, model_slide), item, number, criterion)
    _remove_slide(prs, model_slide)
    filename = scenario_filename(item)
    prs.save(str(folder / filename))
    return filename
