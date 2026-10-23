"""Presentazione PowerPoint della Dashboard di un increment: una slide per
ogni card della pagina (Scope, indicatori, issue chiuse, fasi, grafici), con
gli stessi dati e le stesse regole di calcolo.

Grafica: se esiste TEMPLATE_PATH (un .pptx con la grafica aziendale, oggi
quella Inpeco) la presentazione nasce da quello - schema diapositiva, font,
logo e colori del tema sono i suoi; altrimenti da un layout neutro disegnato
qui. Tabelle e
grafici usano i colori del tema, quindi seguono da soli la grafica del
template. I grafici sono grafici PowerPoint veri (non immagini): i dati
restano modificabili da PowerPoint.
"""

import datetime as dt
import math
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from io import BytesIO

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LABEL_POSITION,
    XL_LEGEND_POSITION,
    XL_MARKER_STYLE,
    XL_TICK_LABEL_POSITION,
)
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from app import models
from app.services.document_generator import TEMPLATES_DIR
from app.services.metrics import compute_dashboard_metrics

# Template con la grafica aziendale: facoltativo, basta copiare qui un .pptx
# (non .potx: python-pptx non lo apre). Le sue slide vengono scartate, si
# usano solo schema e layout.
TEMPLATE_PATH = TEMPLATES_DIR / "presentation_template.pptx"
# Layout del template da usare per copertina e slide di contenuto, per nome
# (minuscolo): prima quelli del template Inpeco, poi quelli standard di
# PowerPoint. Se non ce n'e' nessuno si sceglie in base ai segnaposto.
# Font di tutto cio' che viene disegnato sulle slide del template (testi,
# tabelle, grafici): quello della grafica Inpeco, lo stesso dei segnaposto
# dei suoi layout.
TEMPLATE_FONT = "Montserrat"
COVER_LAYOUT_NAMES = ("cover", "title slide")
CONTENT_LAYOUT_NAMES = ("layout page", "title only")

JIRA_BROWSE_URL = "https://inpeco.atlassian.net/browse/"

# Stessi toni della Dashboard (index.css).
COLOR_TEXT = RGBColor(0x1C, 0x22, 0x2B)
COLOR_MUTED = RGBColor(0x66, 0x70, 0x7D)
COLOR_TILE = RGBColor(0xF0, 0xF2, 0xF5)
COLOR_SUCCESS = RGBColor(0x1A, 0x9C, 0x5C)
COLOR_WARNING = RGBColor(0xC9, 0x7A, 0x12)
COLOR_DANGER = RGBColor(0xD3, 0x40, 0x2F)
COLOR_PLANNED = RGBColor(0xA3, 0xAC, 0xB9)
COLOR_GRID = RGBColor(0xE2, 0xE5, 0xEA)
# Fasi nella timeline: data pianificata e data effettiva, come nel Gantt
# della Dashboard generale (OverviewDashboardPage).
COLOR_PHASE_PLANNED = RGBColor(0x2A, 0x78, 0xD6)
COLOR_PHASE_ACTUAL = RGBColor(0xEB, 0x68, 0x34)
# Gantt del Backlog: stessi colori di BacklogGanttChart (colore = tipo Jira;
# attivita' di servizio a righe, blu il pianificato e arancio l'effettivo).
GANTT_TYPE_COLORS = {
    "Story": RGBColor(0x2E, 0xCC, 0x40),
    "Bug": RGBColor(0xD3, 0x40, 0x2F),
    "Activity": RGBColor(0x2F, 0x6F, 0xED),
}
COLOR_GANTT_OTHER = RGBColor(0x5B, 0x64, 0x72)
COLOR_GANTT_SERVICE_PLANNED = RGBColor(0x2F, 0x6F, 0xED)
COLOR_GANTT_SERVICE_ACTUAL = RGBColor(0xEB, 0x68, 0x34)
COLOR_GANTT_PHASE = RGBColor(0x7C, 0x4D, 0xBD)
MONTH_NAMES = ("gen", "feb", "mar", "apr", "mag", "giu", "lug", "ago", "set", "ott", "nov", "dic")
# Evidenziatore (giallo chiaro) delle voci dello Scope in lavorazione.
COLOR_HIGHLIGHT = "FFE98A"

# Chiave di una issue Jira citata in un testo (es. "PTBSYS-5778").
JIRA_KEY_RE = re.compile(r"\b[A-Z][A-Z0-9]+-\d+\b")

# Oltre questo numero di PBI un grafico previsione vs effettivo diventa
# illeggibile su una slide: si spezza su piu' slide.
PLAN_VS_ACTUAL_POINTS_PER_SLIDE = 30

EXCEL_EPOCH = dt.date(1899, 12, 30)


def _num(value: float, digits: int = 1) -> str:
    """Numero in formato italiano (1.234,5), senza decimali inutili."""
    text = f"{value:,.{digits}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if digits and "," in text:
        text = text.rstrip("0").rstrip(",")
    return text


def _pct(value: float | None) -> str:
    return "—" if value is None else f"{round(value * 100)}%"


def _date(value: dt.date | None) -> str:
    return value.strftime("%d/%m/%Y") if value else "—"


def _working_days_between(start: dt.date | None, end: dt.date | None) -> int | None:
    """Giorni lavorativi (lun-ven) tra due date, estremi inclusi: come
    workingDaysBetween del frontend (colonna Durata (gg) del Backlog)."""
    if start is None or end is None or end < start:
        return None
    return sum(1 for n in range((end - start).days + 1) if (start + dt.timedelta(days=n)).weekday() < 5)


def _spi_status(spi: float | None) -> tuple[str | None, RGBColor]:
    if spi is None:
        return None, COLOR_TEXT
    if spi > 1.05:
        return "Ahead of Schedule", COLOR_SUCCESS
    if spi < 0.95:
        return "Behind Schedule", COLOR_WARNING
    return "On Schedule", COLOR_TEXT


def _period_title(days: int) -> str:
    return {7: "nell'ultima settimana", 14: "nelle ultime 2 settimane", 30: "nell'ultimo mese"}.get(
        days, f"negli ultimi {days} giorni"
    )


# --- Scope: da HTML (editor della Dashboard) a paragrafi ---------------------


@dataclass
class _ScopeParagraph:
    level: int = 0
    marker: str = ""
    heading: bool = False
    runs: list[tuple[str, bool]] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "".join(t for t, _ in self.runs)


class _ScopeParser(HTMLParser):
    """Lo Scope e' l'HTML salvato dall'editor (vedi lib/richText nel
    frontend): qui se ne tengono struttura e grassetti - paragrafi, titoli,
    elenchi annidati (puntati, numerati, di controllo), righe di tabella."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.paragraphs: list[_ScopeParagraph] = []
        self._lists: list[dict] = []
        self._bold = 0
        self._cell_depth = 0
        self._current: _ScopeParagraph | None = None

    def _new_paragraph(self, marker: str = "", heading: bool = False) -> None:
        self._current = _ScopeParagraph(level=max(len(self._lists) - 1, 0), marker=marker, heading=heading)
        self.paragraphs.append(self._current)

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in ("ul", "ol"):
            start = attrs.get("start") or "1"
            self._lists.append({"ordered": tag == "ol", "n": int(start) if start.isdigit() else 1})
        elif tag == "li":
            current_list = self._lists[-1] if self._lists else {"ordered": False, "n": 1}
            if attrs.get("data-type") == "taskItem":
                marker = "☑ " if attrs.get("data-checked") == "true" else "☐ "
            elif current_list["ordered"]:
                marker = f"{current_list['n']}. "
                current_list["n"] += 1
            else:
                marker = "• "
            self._new_paragraph(marker)
        elif tag in ("p", "h1", "h2", "h3", "blockquote", "pre", "tr"):
            # Il <p> dentro una voce d'elenco o una cella non apre un nuovo
            # paragrafo: e' il testo di quella voce/cella.
            if self._cell_depth and tag != "tr":
                return
            if self._current is None or self._current.runs:
                self._new_paragraph(heading=tag in ("h1", "h2", "h3"))
        elif tag in ("td", "th"):
            if self._current is not None and self._current.runs:
                self._current.runs.append(("  |  ", False))
            self._cell_depth += 1
        elif tag in ("strong", "b"):
            self._bold += 1
        elif tag == "br" and self._current is not None:
            self._current.runs.append((" ", False))

    def handle_endtag(self, tag):
        if tag in ("ul", "ol") and self._lists:
            self._lists.pop()
            self._current = None
        elif tag in ("td", "th"):
            self._cell_depth = max(self._cell_depth - 1, 0)
        elif tag in ("strong", "b"):
            self._bold = max(self._bold - 1, 0)
        elif tag in ("h1", "h2", "h3", "tr", "blockquote", "pre"):
            self._current = None

    def handle_data(self, data):
        if not data.strip() and (self._current is None or not self._current.runs):
            return
        if self._current is None:
            self._new_paragraph()
        self._current.runs.append((data, self._bold > 0))


def _scope_paragraphs(scope: str | None) -> list[_ScopeParagraph]:
    if not scope or not scope.strip():
        return []
    if not re.search(r"<[a-z][^>]*>", scope, flags=re.IGNORECASE):
        # Scope scritto prima dell'editor formattato: testo semplice.
        return [_ScopeParagraph(runs=[(line.strip(), False)]) for line in scope.splitlines() if line.strip()]
    parser = _ScopeParser()
    parser.feed(scope)
    return [p for p in parser.paragraphs if p.text.strip()]


# --- Presentazione ------------------------------------------------------------


@dataclass
class _Box:
    left: int
    top: int
    width: int
    height: int


class _Deck:
    """Presentazione in costruzione: nasconde la differenza tra il template
    aziendale (titoli nei segnaposto dei suoi layout) e il layout neutro
    (titoli disegnati qui)."""

    def __init__(self, footer: str) -> None:
        self.footer = footer
        self.cover_layout = None
        self.content_layout = None
        # Font applicato a fine lavoro a tutti i testi delle slide (segnaposto,
        # caselle di testo, tabelle, grafici).
        self.font_name: str | None = None
        # Titoli disegnati qui (non nel segnaposto titolo), per retitle().
        self._heading_runs: dict[int, object] = {}
        if TEMPLATE_PATH.exists():
            self.prs = Presentation(str(TEMPLATE_PATH))
            self._drop_slides()
            layouts = list(self.prs.slide_layouts)

            def placeholder_types(layout):
                return [ph.placeholder_format.type for ph in layout.placeholders]

            def by_name(names):
                return next((lay for name in names for lay in layouts if lay.name.lower() == name), None)

            self.cover_layout = by_name(COVER_LAYOUT_NAMES) or next(
                (lay for lay in layouts if PP_PLACEHOLDER.CENTER_TITLE in placeholder_types(lay)), layouts[0]
            )
            # Layout "solo titolo" (il contenuto lo disegna questa classe): per
            # nome, altrimenti quello col titolo e meno segnaposto.
            with_title = [lay for lay in layouts if PP_PLACEHOLDER.TITLE in placeholder_types(lay)]
            self.content_layout = by_name(CONTENT_LAYOUT_NAMES) or (
                min(with_title, key=lambda lay: len(lay.placeholders)) if with_title else None
            )
            self.font_name = TEMPLATE_FONT
        else:
            self.prs = Presentation()
            self.prs.slide_width = Inches(13.333)
            self.prs.slide_height = Inches(7.5)
        self.blank_layout = min(self.prs.slide_layouts, key=lambda lay: len(lay.placeholders))
        self.width = self.prs.slide_width
        self.height = self.prs.slide_height

    def _drop_slides(self) -> None:
        slide_ids = self.prs.slides._sldIdLst
        for slide_id in list(slide_ids):
            self.prs.part.drop_rel(slide_id.rId)
            slide_ids.remove(slide_id)

    @staticmethod
    def _remove_empty_placeholders(slide) -> None:
        for ph in list(slide.placeholders):
            if not (ph.has_text_frame and ph.text_frame.text.strip()):
                ph._element.getparent().remove(ph._element)

    def add_cover(self, title: str, lines: list[str]) -> None:
        """Copertina: il titolo e, sotto, le righe date (una per segnaposto
        di testo del layout di copertina, dall'alto in basso)."""
        if self.cover_layout is not None:
            slide = self.prs.slides.add_slide(self.cover_layout)
            if slide.shapes.title is not None:
                slide.shapes.title.text = title
            bodies = sorted(
                (ph for ph in slide.placeholders if ph != slide.shapes.title and ph.has_text_frame),
                key=lambda ph: ph.top or 0,
            )
            for index, body in enumerate(bodies[: len(lines)]):
                # Le righe in piu' rispetto ai segnaposto finiscono nell'ultimo.
                last = index == min(len(bodies), len(lines)) - 1
                body.text = "   ·   ".join(lines[index:]) if last else lines[index]
            self._remove_empty_placeholders(slide)
            return
        slide = self.prs.slides.add_slide(self.blank_layout)
        self._accent_bar(slide, 0, 0, self.width, Inches(0.35))
        margin = Inches(0.9)
        _text(slide, _Box(margin, Inches(2.6), self.width - 2 * margin, Inches(1.4)), title, size=40, bold=True)
        _text(
            slide,
            _Box(margin, Inches(4.1), self.width - 2 * margin, Inches(1.0)),
            "   ·   ".join(lines),
            size=20,
            color=COLOR_MUTED,
        )

    def add_slide(self, title: str):
        """Nuova slide col titolo dato: restituisce la slide e l'area libera
        sotto il titolo, in cui disegnare il contenuto."""
        margin = Inches(0.6)
        if self.content_layout is not None:
            slide = self.prs.slides.add_slide(self.content_layout)
            title_shape = slide.shapes.title
            if (title_shape.top or 0) > self.height / 2:
                # Template (come quello Inpeco) in cui il segnaposto "titolo"
                # e' una riga piccola a pie' di pagina: li' va il nome
                # dell'increment, e il titolo della slide si disegna in alto
                # a sinistra, nel colore del tema, accanto al logo.
                title_shape.text = self.footer
                self._remove_empty_placeholders(slide)
                margin = Inches(0.48)
                heading = _text_frame(slide, _Box(margin, Inches(0.4), self.width - margin - Inches(1.8), Inches(0.7)))
                run = _run(heading.paragraphs[0], title, size=28)
                run.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
                self._heading_runs[slide.slide_id] = run
                top = Inches(1.3)
                return slide, _Box(margin, top, self.width - 2 * margin, title_shape.top - top - Inches(0.2))
            title_shape.text = title
            self._remove_empty_placeholders(slide)
            left = title_shape.left if title_shape.left is not None else margin
            width = title_shape.width if title_shape.width is not None else self.width - 2 * margin
            top = (title_shape.top or 0) + (title_shape.height or Inches(1.1)) + Inches(0.15)
            # In basso resta libera la fascia di logo/pie' di pagina del template.
            return slide, _Box(left, top, width, self.height - top - Inches(0.9))

        slide = self.prs.slides.add_slide(self.blank_layout)
        heading = _text_frame(slide, _Box(margin, Inches(0.4), self.width - 2 * margin, Inches(0.7)))
        self._heading_runs[slide.slide_id] = _run(heading.paragraphs[0], title, size=28, bold=True)
        self._accent_bar(slide, margin, Inches(1.15), Inches(1.2), Inches(0.06))
        _text(
            slide,
            _Box(margin, self.height - Inches(0.5), self.width - 2 * margin, Inches(0.3)),
            self.footer,
            size=10,
            color=COLOR_MUTED,
        )
        top = Inches(1.45)
        return slide, _Box(margin, top, self.width - 2 * margin, self.height - top - Inches(0.7))

    def retitle(self, slide, title: str) -> None:
        """Cambia il titolo di una slide creata da add_slide (es. per
        aggiungere "1/3" quando il contenuto si scopre lungo piu' slide)."""
        run = self._heading_runs.get(slide.slide_id)
        if run is not None:
            run.text = title
        else:
            slide.shapes.title.text = title

    def _apply_font(self) -> None:
        font_name = self.font_name
        if not font_name:
            return

        def style(text_frame) -> None:
            for paragraph in text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = font_name

        for slide in self.prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    style(shape.text_frame)
                elif getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            style(cell.text_frame)
                elif getattr(shape, "has_chart", False) and shape.has_chart:
                    chart = shape.chart
                    chart.font.name = font_name
                    if chart.value_axis.has_title:
                        style(chart.value_axis.axis_title.text_frame)

    @staticmethod
    def _accent_bar(slide, left, top, width, height) -> None:
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        bar.fill.solid()
        bar.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        bar.line.fill.background()
        bar.shadow.inherit = False

    def to_bytes(self) -> bytes:
        self._apply_font()
        buffer = BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()


def _run(paragraph, text: str, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None):
    run = paragraph.add_run()
    run.text = text
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    return run


def _highlight(run, color: str = COLOR_HIGHLIGHT) -> None:
    """Evidenziatore sul testo (python-pptx non lo espone: si scrive
    l'elemento a:highlight, che va prima di ogni altro figlio gia' presente
    tranne contorno e riempimento)."""
    r_pr = run._r.get_or_add_rPr()
    highlight = r_pr.makeelement(qn("a:highlight"), {})
    highlight.append(highlight.makeelement(qn("a:srgbClr"), {"val": color}))
    before = [child for child in r_pr if child.tag in (qn("a:ln"), qn("a:solidFill"), qn("a:effectLst"))]
    r_pr.insert(len(before), highlight)


def _text_frame(slide, box: _Box):
    frame = slide.shapes.add_textbox(box.left, box.top, box.width, box.height).text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    return frame


def _text(slide, box: _Box, text: str, size: float = 14, bold: bool = False, color: RGBColor | None = None):
    frame = _text_frame(slide, box)
    _run(frame.paragraphs[0], text, size=size, bold=bold, color=color)
    return frame


def _add_table(
    slide,
    box: _Box,
    headers: list[str],
    rows: list[list[str]],
    col_ratios: list[float],
    bold_row: int | None = None,
    link_col: int | None = None,
    font_size: float = 11,
) -> None:
    """Tabella con lo stile di default (intestazione nel colore del tema).
    link_col: colonna con chiavi Jira, rese link alla issue."""
    row_height = Inches(0.36)
    shape = slide.shapes.add_table(len(rows) + 1, len(headers), box.left, box.top, box.width, row_height * (len(rows) + 1))
    table = shape.table
    total = sum(col_ratios)
    for index, ratio in enumerate(col_ratios):
        table.columns[index].width = Emu(int(box.width * ratio / total))
    for r, values in enumerate([headers, *rows]):
        table.rows[r].height = row_height
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            run = _run(cell.text_frame.paragraphs[0], value, size=font_size, bold=True if r == 0 or r - 1 == bold_row else None)
            if r > 0 and c == link_col and value:
                run.hyperlink.address = f"{JIRA_BROWSE_URL}{value}"


def _style_chart(chart, legend: bool = True, value_axis_title: str | None = None) -> None:
    chart.font.size = Pt(11)
    # Il titolo e' gia' quello della slide.
    chart.has_title = False
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    # Griglia chiara: di default e' nera e copre i dati.
    for axis in (chart.category_axis, chart.value_axis):
        if axis.has_major_gridlines:
            axis.major_gridlines.format.line.color.rgb = COLOR_GRID
    if value_axis_title:
        chart.value_axis.has_title = True
        title = chart.value_axis.axis_title.text_frame.paragraphs[0]
        _run(title, value_axis_title, size=11, bold=False, color=COLOR_MUTED)


def _scope_slides(deck: _Deck, project: models.Project) -> None:
    paragraphs = _scope_paragraphs(project.scope)
    if not paragraphs:
        slide, box = deck.add_slide("Scope")
        _text(slide, box, "Nessuno scope definito.", color=COLOR_MUTED)
        return

    slide, box = deck.add_slide("Scope")
    in_progress_keys = {i.jira_key for i in project.backlog_items if i.status == "In Progress"}

    def paginate(font_size: float) -> list[list[_ScopeParagraph]]:
        # Stima di quante righe occupa ogni paragrafo, per spezzare uno
        # Scope lungo su piu' slide invece di farlo uscire dal fondo.
        chars_per_line = max(int(box.width / Pt(font_size) * 1.6), 20)
        lines_per_slide = max(int(box.height / Pt(font_size * 1.2)), 4)
        pages: list[list[_ScopeParagraph]] = [[]]
        used = 0
        for paragraph in paragraphs:
            width = max(chars_per_line - paragraph.level * 5, 20)
            lines = math.ceil(len(paragraph.marker + paragraph.text) / width) or 1
            if pages[-1] and used + lines > lines_per_slide:
                pages.append([])
                used = 0
            pages[-1].append(paragraph)
            used += lines
        return pages

    # Il carattere piu' grande con cui lo Scope sta in una slide sola; se non
    # ci sta nemmeno col piu' piccolo, si va su piu' slide.
    for font_size in (16, 14, 13, 12):
        pages = paginate(font_size)
        if len(pages) == 1:
            break

    for index, page in enumerate(pages):
        if index > 0:
            slide, box = deck.add_slide(f"Scope ({index + 1}/{len(pages)})")
        elif len(pages) > 1:
            deck.retitle(slide, f"Scope (1/{len(pages)})")
        frame = _text_frame(slide, box)
        for n, paragraph in enumerate(page):
            p = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
            # Il rientro degli elenchi annidati: spazi davanti al marcatore,
            # indipendenti dagli stili di elenco del template.
            _run(p, "      " * paragraph.level + paragraph.marker, size=font_size)
            # Le voci che citano una issue in lavorazione ora (es.
            # "... - PTBSYS-5778") sono evidenziate, con un razzo a destra.
            in_progress = any(key in in_progress_keys for key in JIRA_KEY_RE.findall(paragraph.text))
            for text, bold in paragraph.runs:
                run = _run(p, text, size=font_size, bold=True if (bold or paragraph.heading) else None)
                if in_progress:
                    _highlight(run)
            if in_progress:
                _run(p, "  \U0001F680", size=font_size)


def _kpi_slide(deck: _Deck, project: models.Project) -> None:
    metrics = compute_dashboard_metrics(project)
    slide, box = deck.add_slide("Indicatori")
    snapshot_note = f"da snapshot del {_date(metrics.last_snapshot_date)}"
    status_label, status_color = _spi_status(metrics.spi)

    # (valore, colore valore, etichetta, note[(testo, colore)])
    tiles: list[tuple[str, RGBColor, str, list[tuple[str, RGBColor]]]] = [
        (
            _pct(metrics.percent_complete),
            COLOR_TEXT,
            f"Completamento backlog ({metrics.backlog_done}/{metrics.backlog_in_scope})",
            [(snapshot_note, COLOR_MUTED)] if metrics.completion_source == "snapshot" else [],
        ),
        (
            f"{_num(metrics.logged_hours_total)} h",
            COLOR_TEXT,
            "Ore usate",
            (
                [
                    (
                        f"{_pct(metrics.percent_budget_used)} del budget di {_num(metrics.budget_hours_total)} h",
                        COLOR_DANGER if metrics.percent_budget_used > 1 else COLOR_TEXT,
                    )
                ]
                if metrics.percent_budget_used is not None
                else []
            )
            + ([(snapshot_note, COLOR_MUTED)] if metrics.logged_hours_source == "snapshot" else [])
            + ([("somma dei progetti collegati", COLOR_MUTED)] if metrics.logged_hours_source == "progetti" else []),
        ),
        (
            f"{metrics.spi:.2f}".replace(".", ",") if metrics.spi is not None else "—",
            status_color,
            "SPI (avanzamento / tempo trascorso)",
            [],
        ),
        (status_label or "—", status_color, "Status", []),
    ]

    gap = Inches(0.2)
    tile_width = int((box.width - gap * (len(tiles) - 1)) / len(tiles))
    tile_height = min(Inches(1.9), box.height)
    for index, (value, value_color, label, notes) in enumerate(tiles):
        tile = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, box.left + index * (tile_width + gap), box.top, tile_width, tile_height
        )
        tile.adjustments[0] = 0.06
        tile.fill.solid()
        tile.fill.fore_color.rgb = COLOR_TILE
        tile.line.fill.background()
        tile.shadow.inherit = False
        frame = tile.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = frame.margin_right = Inches(0.18)
        frame.margin_top = Inches(0.2)
        # Lo Status e' un testo, non un numero: piu' piccolo per stare nel riquadro.
        _run(frame.paragraphs[0], value, size=20 if index == len(tiles) - 1 else 30, bold=True, color=value_color)
        _run(frame.add_paragraph(), label, size=12, color=COLOR_MUTED)
        for note, note_color in notes:
            _run(frame.add_paragraph(), note, size=10, color=note_color)
        for paragraph in frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT

    details = [
        f"Stato increment: {project.status}",
        f"Inizio: {_date(project.start_date)}",
        f"Code freeze: {_date(project.code_freeze_date)}",
    ]
    if project.release_date is not None:
        details.append(f"Rilasciato il {_date(project.release_date)}")
    if metrics.percent_time_elapsed is not None:
        details.append(f"Tempo trascorso: {_pct(metrics.percent_time_elapsed)}")
    details_top = box.top + tile_height + Inches(0.35)
    if details_top + Inches(0.4) <= box.top + box.height:
        _text(slide, _Box(box.left, details_top, box.width, Inches(0.4)), "   ·   ".join(details), size=14, color=COLOR_MUTED)


def _recently_closed_slides(deck: _Deck, project: models.Project, days: int) -> None:
    cutoff = dt.date.today() - dt.timedelta(days=days)
    closed = sorted(
        (
            i
            for i in project.backlog_items
            if i.is_in_plan and i.status == "Done" and i.actual_finish is not None and i.actual_finish >= cutoff
        ),
        key=lambda i: (-i.actual_finish.toordinal(), i.priority_order),
    )
    # Sotto le chiuse, quelle in lavorazione ora (stato In Progress, vedi
    # BacklogItem.status) e quelle ancora da iniziare, nell'ordine del
    # Backlog. Tutte e tre con la stessa regola del piano: lavoro del code
    # freeze e attivita' di servizio, come il Gantt del Backlog.
    in_progress = sorted(
        (i for i in project.backlog_items if i.is_in_plan and i.status == "In Progress"),
        key=lambda i: i.priority_order,
    )
    to_do = sorted(
        (i for i in project.backlog_items if i.is_in_plan and i.status == "To Do"),
        key=lambda i: i.priority_order,
    )
    # (titolo della sezione, intestazione dell'ultima colonna, righe, testo se vuota)
    sections = [
        (
            f"Chiuse {_period_title(days)} ({len(closed)})",
            "Chiusa il",
            [[i.jira_key, i.issue_type or "-", i.summary or "-", _date(i.actual_finish)] for i in closed],
            f"Nessuna issue in scope chiusa negli ultimi {days} giorni.",
        ),
        (
            f"In progress ({len(in_progress)})",
            "Iniziata il",
            [[i.jira_key, i.issue_type or "-", i.summary or "-", _date(i.actual_start)] for i in in_progress],
            "Nessuna issue in scope in progress.",
        ),
        (
            f"Da fare ({len(to_do)})",
            "Stato Jira",
            [[i.jira_key, i.issue_type or "-", i.summary or "-", i.jira_status or "To Do"] for i in to_do],
            "Nessuna issue in scope ancora da iniziare.",
        ),
    ]

    title = "Issue chiuse, in progress e da fare"
    heading_height = Inches(0.4)
    header_height = Inches(0.36)

    def row_height(row: list[str]) -> int:
        # Le Summary lunghe vanno a capo e alzano la riga.
        return Inches(0.36) * min(math.ceil(max(len(row[2]), 1) / 95), 4)

    slide, box = deck.add_slide(title)
    slides = [slide]
    top = box.top
    bottom = box.top + box.height
    for heading, date_header, rows, empty_message in sections:
        continued = False
        while True:
            # Quante righe della sezione stanno nello spazio rimasto: se non
            # ce ne sta nemmeno una si passa a una nuova slide.
            available = bottom - top - heading_height - header_height
            fitting = 0
            used = 0
            for row in rows:
                if used + row_height(row) > available:
                    break
                used += row_height(row)
                fitting += 1
            if (rows and fitting == 0) or (not rows and available < 0):
                slide, box = deck.add_slide(title)
                slides.append(slide)
                top = box.top
                continue
            _text(
                slide,
                _Box(box.left, top, box.width, heading_height),
                heading + (" - segue" if continued else ""),
                size=14,
                bold=True,
            )
            top += heading_height
            if not rows:
                _text(slide, _Box(box.left, top, box.width, heading_height), empty_message, size=12, color=COLOR_MUTED)
                top += heading_height + Inches(0.2)
                break
            _add_table(
                slide,
                _Box(box.left, top, box.width, 0),
                ["ID", "Tipo", "Summary", date_header],
                rows[:fitting],
                col_ratios=[1.5, 1.1, 7.4, 1.3],
                link_col=0,
            )
            top += header_height + used + Inches(0.3)
            rows = rows[fitting:]
            if not rows:
                break
            continued = True
            slide, box = deck.add_slide(title)
            slides.append(slide)
            top = box.top
    if len(slides) > 1:
        for index, page in enumerate(slides):
            deck.retitle(page, f"{title} - {index + 1}/{len(slides)}")


def _phases_slide(deck: _Deck, project: models.Project) -> None:
    slide, box = deck.add_slide("Fasi increment")
    phases = list(project.phases)
    if not phases:
        _text(slide, box, "Nessuna fase definita.", color=COLOR_MUTED)
        return
    current = next((n for n, p in enumerate(phases) if p.name == project.status), None)
    _add_table(
        slide,
        box,
        ["Fase", "Pianificata", "Effettiva", "Note"],
        [
            [
                p.name + ("  (fase corrente)" if n == current else ""),
                _date(p.planned_date) if p.planned_date else "",
                _date(p.actual_date) if p.actual_date else "",
                p.notes or "",
            ]
            for n, p in enumerate(phases)
        ],
        col_ratios=[3.2, 1.6, 1.6, 5],
        bold_row=current,
        font_size=12,
    )


def _month_start(day: dt.date, months: int = 0) -> dt.date:
    """Primo giorno del mese di day, spostato di months mesi."""
    index = day.year * 12 + day.month - 1 + months
    return dt.date(index // 12, index % 12 + 1, 1)


def _timeline_slide(deck: _Deck, project: models.Project) -> None:
    """Gantt delle fasi: una riga per fase, con le sue date e, sull'asse del
    tempo, un rombo vuoto alla data pianificata e uno pieno a quella
    effettiva (come nel Gantt della Dashboard generale). In cima la barra
    dell'increment, dall'inizio al Planned finish; se l'increment e' stato
    rilasciato, una linea verde alla data di rilascio."""
    slide, box = deck.add_slide("Timeline increment")
    phases = list(project.phases)
    start, finish, release = project.start_date, project.planned_finish_date, project.release_date
    days = [d for d in (start, finish, release) if d is not None]
    days += [d for p in phases for d in (p.planned_date, p.actual_date) if d is not None]
    if not days:
        _text(
            slide,
            box,
            "Nessuna data da mostrare: servono le date delle fasi o le date di inizio e Planned finish dell'increment.",
            color=COLOR_MUTED,
        )
        return

    # L'asse va a mesi interi, dal mese della prima data a quello dell'ultima.
    axis_start = _month_start(min(days))
    axis_end = _month_start(max(days), 1)
    span = (axis_end - axis_start).days
    month_count = (axis_end.year - axis_start.year) * 12 + axis_end.month - axis_start.month

    date_width = Inches(1.15)
    name_width = int(box.width * 0.2)
    chart_left = box.left + name_width + 2 * date_width + Inches(0.2)
    # A destra resta mezzo rombo di margine, per una data a fine asse.
    chart_width = box.left + box.width - chart_left - Inches(0.15)
    header_height = Inches(0.4)
    legend_height = Inches(0.45)
    has_increment_row = start is not None or finish is not None
    row_count = len(phases) + (1 if has_increment_row else 0)
    row_height = min(Inches(0.5), int((box.height - header_height - legend_height) / row_count))
    font_size = 12 if row_height >= Inches(0.34) else 10 if row_height >= Inches(0.26) else 8
    rows_top = box.top + header_height
    rows_bottom = rows_top + row_height * row_count

    def x_of(day: dt.date) -> int:
        return chart_left + int(chart_width * (day - axis_start).days / span)

    def line(x1: int, y1: int, x2: int, y2: int, color: RGBColor, width: float = 0.75, dashed: bool = False) -> None:
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        connector.line.color.rgb = color
        connector.line.width = Pt(width)
        if dashed:
            connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    def cell(left: int, top: int, width: int, height: int, text: str, **style) -> None:
        _text(slide, _Box(left, top, width, height), text, **style).vertical_anchor = MSO_ANCHOR.MIDDLE

    def diamond(x: int, y: int, size: int, color: RGBColor, filled: bool) -> None:
        mark = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, x - size // 2, y - size // 2, size, size)
        mark.fill.solid()
        mark.fill.fore_color.rgb = color if filled else RGBColor(0xFF, 0xFF, 0xFF)
        mark.line.color.rgb = color
        mark.line.width = Pt(1.5)
        mark.shadow.inherit = False

    # Intestazione: colonne delle date nei colori dei rispettivi rombi, e i
    # mesi sull'asse (uno ogni tanti quando sono troppi per starci tutti).
    planned_left = box.left + name_width
    actual_left = planned_left + date_width
    cell(box.left, box.top, name_width, header_height, "Fase", size=11, bold=True, color=COLOR_MUTED)
    cell(planned_left, box.top, date_width, header_height, "Pianificata", size=11, bold=True, color=COLOR_PHASE_PLANNED)
    cell(actual_left, box.top, date_width, header_height, "Effettiva", size=11, bold=True, color=COLOR_PHASE_ACTUAL)
    month_step = math.ceil(month_count / 12)
    for n in range(0, month_count + 1, month_step):
        month = _month_start(axis_start, n)
        x = x_of(min(month, axis_end))
        line(x, rows_top, x, rows_bottom, COLOR_GRID)
        if n < month_count:
            label_width = x_of(min(_month_start(month, month_step), axis_end)) - x
            cell(
                x + Inches(0.04),
                box.top,
                max(label_width, Inches(0.7)),
                header_height,
                f"{MONTH_NAMES[month.month - 1]} {month:%y}",
                size=10,
                color=COLOR_MUTED,
            )
    for n in range(row_count + 1):
        y = rows_top + n * row_height
        line(box.left, y, box.left + box.width, y, COLOR_GRID)

    today = dt.date.today()
    show_today = axis_start <= today <= axis_end
    if show_today:
        line(x_of(today), rows_top, x_of(today), rows_bottom, COLOR_DANGER, width=1.25, dashed=True)
    if release is not None:
        line(x_of(release), rows_top, x_of(release), rows_bottom, COLOR_SUCCESS, width=2.25)

    mark_size = min(Inches(0.2), int(row_height * 0.62))
    top = rows_top
    if has_increment_row:
        cell(box.left, top, name_width, row_height, "Increment", size=font_size, bold=True)
        cell(planned_left, top, 2 * date_width, row_height, f"{_date(start)} – {_date(finish)}", size=font_size)
        if start is not None and finish is not None and finish >= start:
            bar_height = int(row_height * 0.4)
            bar = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                x_of(start),
                top + (row_height - bar_height) // 2,
                max(x_of(finish) - x_of(start), Inches(0.03)),
                bar_height,
            )
            bar.fill.solid()
            bar.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
            bar.line.fill.background()
            bar.shadow.inherit = False
        top += row_height
    for phase in phases:
        is_current = phase.name == project.status
        middle = top + row_height // 2
        cell(box.left, top, name_width, row_height, phase.name, size=font_size, bold=is_current)
        cell(planned_left, top, date_width, row_height, _date(phase.planned_date), size=font_size, bold=is_current)
        cell(actual_left, top, date_width, row_height, _date(phase.actual_date), size=font_size, bold=is_current)
        if phase.planned_date is not None and phase.actual_date is not None and phase.planned_date != phase.actual_date:
            # Il tratto tra le due date: lo scostamento della fase.
            line(x_of(phase.planned_date), middle, x_of(phase.actual_date), middle, COLOR_PLANNED, width=1.5)
        if phase.planned_date is not None:
            diamond(x_of(phase.planned_date), middle, mark_size, COLOR_PHASE_PLANNED, filled=False)
        if phase.actual_date is not None:
            # Piu' piccolo: se cade nello stesso giorno della data pianificata
            # resta dentro il rombo vuoto, e si vedono entrambi.
            diamond(x_of(phase.actual_date), middle, int(mark_size * 0.7), COLOR_PHASE_ACTUAL, filled=True)
        top += row_height

    # Legenda, sotto le righe.
    legend_top = rows_bottom + Inches(0.1)
    legend_middle = legend_top + Inches(0.15)
    x = box.left + Inches(0.1)
    legend = [("Data pianificata", COLOR_PHASE_PLANNED, False), ("Data effettiva", COLOR_PHASE_ACTUAL, True)]
    for label, color, filled in legend:
        diamond(x, legend_middle, Inches(0.16) if not filled else Inches(0.12), color, filled)
        cell(x + Inches(0.18), legend_top, Inches(1.7), Inches(0.3), label, size=10, color=COLOR_MUTED)
        x += Inches(1.9)
    if show_today:
        line(x - Inches(0.1), legend_middle, x + Inches(0.2), legend_middle, COLOR_DANGER, width=1.25, dashed=True)
        cell(x + Inches(0.3), legend_top, Inches(1.8), Inches(0.3), f"Oggi ({_date(today)})", size=10, color=COLOR_MUTED)
        x += Inches(2.1)
    if release is not None:
        line(x - Inches(0.1), legend_middle, x + Inches(0.2), legend_middle, COLOR_SUCCESS, width=2.25)
        cell(x + Inches(0.3), legend_top, Inches(2.0), Inches(0.3), f"Rilascio ({_date(release)})", size=10, color=COLOR_MUTED)
        x += Inches(2.3)
    if any(p.name == project.status for p in phases):
        cell(x, legend_top, Inches(2.5), Inches(0.3), "In grassetto la fase corrente", size=10, color=COLOR_MUTED)


def _backlog_gantt_slides(deck: _Deck, project: models.Project) -> None:
    """Gantt del Backlog, come la vista Gantt della pagina Backlog coi filtri
    di default: item del piano (code freeze e attivita' di servizio) non
    ancora Done, senza i Bug CVE, nell'ordine del Backlog. Per ogni item la
    barra pianificata (chiara, col bordo) sopra quella effettiva (piena, o
    sfumata fino a oggi se e' in corso), nel colore del tipo Jira; le
    attivita' di servizio a righe oblique, blu/arancio. Linee verticali per
    oggi, Code Freeze e fasi. Tutto in una slide: con tanti item le righe si
    stringono, mentre le date dell'asse restano grandi."""
    title = "Gantt Backlog"
    today = dt.date.today()
    items = sorted(
        (
            i
            for i in project.backlog_items
            if i.is_in_plan and i.status != "Done" and not models.is_cve_bug(i.issue_type, i.summary)
        ),
        key=lambda i: i.priority_order,
    )
    def bar_end(item) -> dt.date | None:
        """Fine della barra effettiva: la fine effettiva, oppure oggi se
        l'item e' in corso; None se la barra non si disegna."""
        if item.actual_start is None:
            return None
        if item.actual_finish is not None:
            return item.actual_finish
        return max(item.actual_start, today) if item.status == "In Progress" else None

    def bars(item) -> list[tuple[dt.date, dt.date]]:
        result = []
        if item.planned_start and item.expected_finish and item.expected_finish >= item.planned_start:
            result.append((item.planned_start, item.expected_finish))
        end = bar_end(item)
        if end is not None and end >= item.actual_start:
            result.append((item.actual_start, end))
        return result

    # Solo gli item con almeno una barra, e l'asse solo sulle barre
    # disegnate: un vecchio inizio effettivo di un item non in corso (senza
    # barra) allargherebbe l'asse di anni.
    items = [i for i in items if bars(i)]
    if not items:
        slide, box = deck.add_slide(title)
        _text(
            slide,
            box,
            "Nessun item da mostrare: servono item del piano non ancora chiusi, con le date pianificate o "
            "almeno l'inizio effettivo.",
            color=COLOR_MUTED,
        )
        return

    # Le fasi come nel Gantt del Backlog: alla data effettiva se c'e' gia',
    # altrimenti a quella pianificata.
    marks = [(p.actual_date or p.planned_date, p.name, COLOR_GANTT_PHASE) for p in project.phases]
    if project.code_freeze_date is not None:
        marks.append((project.code_freeze_date, "Code Freeze", COLOR_SUCCESS))
    marks = [m for m in marks if m[0] is not None]

    item_days = [d for i in items for bar in bars(i) for d in bar]
    first, last = min(item_days + [today]), max(item_days + [today])
    # Item ancora "In Progress" su Jira da anni allargherebbero l'asse fino a
    # schiacciare tutto il resto: l'asse parte al massimo 3 mesi prima del
    # primo inizio pianificato (o di oggi), e le barre che iniziano prima
    # partono dal bordo sinistro, con scritto da quando (vedi sotto).
    planned_starts = [i.planned_start for i in items if i.planned_start]
    first = max(first, min(planned_starts + [today]) - dt.timedelta(days=90))
    # Asse a settimane (dal lunedi') se il periodo e' breve, altrimenti a mesi.
    weekly = (last - first).days <= 120
    if weekly:
        axis_start = first - dt.timedelta(days=first.weekday())
        axis_end = last + dt.timedelta(days=7 - last.weekday())
    else:
        axis_start, axis_end = _month_start(first), _month_start(last, 1)
    span = (axis_end - axis_start).days
    marks = sorted((m for m in marks if axis_start <= m[0] <= axis_end), key=lambda m: m[0])

    mark_lane_height = Inches(0.21)
    slide, box = deck.add_slide(title)
    label_width = int(box.width * 0.24)
    chart_left = box.left + label_width
    chart_width = box.left + box.width - chart_left

    def x_of(day: dt.date) -> int:
        return chart_left + int(chart_width * (day - axis_start).days / span)

    # Nomi di Code Freeze e fasi su piu' righe quando si toccherebbero
    # (come nel Gantt del Backlog): ognuno nella prima riga libera.
    lane_ends: list[int] = []
    mark_lanes = []
    for day, name, color in marks:
        text_width = Inches(0.075) * len(name)
        left = x_of(day) + Inches(0.03)
        # Vicino al bordo destro il nome va a sinistra della linea.
        if left + text_width > box.left + box.width:
            left = x_of(day) - Inches(0.03) - text_width
        lane = next((n for n, end in enumerate(lane_ends) if end < left), len(lane_ends))
        if lane == len(lane_ends):
            lane_ends.append(0)
        lane_ends[lane] = left + text_width
        mark_lanes.append((left, text_width, name, color, lane))
    axis_label_height = Inches(0.32)
    header_height = axis_label_height + Inches(0.04) + mark_lane_height * len(lane_ends)
    legend_height = Inches(0.45)
    row_height = min(Inches(0.42), int((box.height - header_height - legend_height) / len(items)))
    # Testo delle righe proporzionato all'altezza della riga.
    row_font = max(5.5, min(9.0, round(row_height / 12700 * 0.62 * 2) / 2))
    rows_top = box.top + header_height
    rows_bottom = rows_top + row_height * len(items)

    def line(x1: int, y1: int, x2: int, y2: int, color: RGBColor, width: float = 0.75, dashed: bool = False):
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        connector.line.color.rgb = color
        connector.line.width = Pt(width)
        if dashed:
            connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    def cell(left: int, top: int, width: int, height: int, text: str, **style):
        frame = _text(slide, _Box(left, top, width, height), text, **style)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        return frame

    # Intestazione: tacche dell'asse (settimane o mesi) e, sotto, i nomi
    # di Code Freeze e fasi sopra le rispettive linee.
    if weekly:
        ticks = [axis_start + dt.timedelta(days=7 * n) for n in range(span // 7 + 1)]
        step = math.ceil(len(ticks) / 14)
        labels = [(d, f"{d:%d}/{d:%m}") for d in ticks[:-1][::step]]
        grid = ticks
    else:
        month_count = (axis_end.year - axis_start.year) * 12 + axis_end.month - axis_start.month
        grid = [_month_start(axis_start, n) for n in range(month_count + 1)]
        step = math.ceil(month_count / 12)
        labels = [(d, f"{MONTH_NAMES[d.month - 1]} {d:%y}") for d in grid[:-1][::step]]
    for day in grid:
        line(x_of(day), rows_top, x_of(day), rows_bottom, COLOR_GRID)
    label_slot = int(chart_width * (7 * step if weekly else 30 * step) / span)
    for day, text in labels:
        cell(x_of(day) + Inches(0.03), box.top, max(label_slot, Inches(0.7)), axis_label_height, text, size=13, bold=True, color=COLOR_MUTED).word_wrap = False
    for left, text_width, name, color, lane in mark_lanes:
        top = box.top + axis_label_height + Inches(0.04) + lane * mark_lane_height
        cell(left, top, text_width, mark_lane_height, name, size=9, bold=True, color=color).word_wrap = False
    for n in range(len(items) + 1):
        y = rows_top + n * row_height
        line(box.left, y, box.left + box.width, y, COLOR_GRID)

    bar_height = int(row_height * 0.3)
    for row, item in enumerate(items):
        top = rows_top + row * row_height
        frame = cell(box.left, top, label_width - Inches(0.1), row_height, "", size=row_font)
        frame.word_wrap = False
        paragraph = frame.paragraphs[0]
        key_run = _run(paragraph, item.jira_key, size=row_font, bold=True)
        key_run.hyperlink.address = f"{JIRA_BROWSE_URL}{item.jira_key}"
        summary = item.summary or ""
        # Il summary quanto ci sta nella colonna, al corpo del testo.
        max_chars = int(label_width / 12700 / (row_font * 0.5)) - len(item.jira_key) - 4
        if len(summary) > max_chars:
            summary = summary[: max(max_chars - 1, 0)] + "…"
        _run(paragraph, "  " + summary, size=row_font, color=COLOR_MUTED)

        color = GANTT_TYPE_COLORS.get(item.issue_type or "", COLOR_GANTT_OTHER)

        def bar(start: dt.date, end: dt.date, y: int):
            start = max(start, axis_start)
            left = x_of(start)
            # Estremi inclusi: la barra copre anche il giorno di fine.
            width = max(x_of(end + dt.timedelta(days=1)) - left, Inches(0.04))
            return slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, y, width, bar_height)

        if item.planned_start and item.expected_finish and item.expected_finish >= item.planned_start:
            shape = bar(item.planned_start, item.expected_finish, top + int(row_height * 0.14))
            if item.is_activity:
                _striped(shape, COLOR_GANTT_SERVICE_PLANNED)
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb = _tint(color, 0.3)
                shape.line.color.rgb = color
                shape.line.width = Pt(0.75)
            shape.shadow.inherit = False
            if item.is_activity:
                shape.line.fill.background()
        end = bar_end(item)
        if item.actual_start is not None and end is not None and end >= item.actual_start:
            shape = bar(item.actual_start, end, top + int(row_height * 0.56))
            in_progress = item.actual_finish is None
            if in_progress:
                _fading(shape, COLOR_DANGER if item.is_activity else color)
            elif item.is_activity:
                _striped(shape, COLOR_GANTT_SERVICE_ACTUAL)
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb = color
            shape.line.fill.background()
            shape.shadow.inherit = False
            if item.actual_start < axis_start:
                # Iniziata prima dell'inizio dell'asse: la data vera, sopra
                # la barra (sulla barra, sottile, non si leggerebbe).
                note = cell(
                    x_of(axis_start) + Inches(0.04),
                    top,
                    Inches(1.4),
                    int(row_height * 0.55),
                    f"◀ dal {_date(item.actual_start)}",
                    size=max(5.5, row_font - 1),
                    bold=True,
                    color=COLOR_MUTED,
                )
                note.word_wrap = False

    for day, _, color in marks:
        line(x_of(day), rows_top, x_of(day), rows_bottom, color, width=1.25, dashed=True)
    if axis_start <= today <= axis_end:
        line(x_of(today), rows_top, x_of(today), rows_bottom, COLOR_DANGER, width=1.5)

    # Legenda, sotto le righe.
    legend_top = rows_bottom + Inches(0.12)
    swatch = Inches(0.16)
    x = box.left

    def legend_entry(label: str, fill, width: int = Inches(1.1)) -> None:
        nonlocal x
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, legend_top + Inches(0.07), swatch, swatch)
        fill(shape)
        shape.shadow.inherit = False
        cell(x + swatch + Inches(0.06), legend_top, width, Inches(0.3), label, size=9, color=COLOR_MUTED)
        x += swatch + Inches(0.06) + width

    def solid(color: RGBColor):
        def apply(shape):
            shape.fill.solid()
            shape.fill.fore_color.rgb = color
            shape.line.fill.background()

        return apply

    def outlined(shape):
        shape.fill.solid()
        shape.fill.fore_color.rgb = _tint(COLOR_GANTT_OTHER, 0.3)
        shape.line.color.rgb = COLOR_GANTT_OTHER
        shape.line.width = Pt(0.75)

    for type_name, color in GANTT_TYPE_COLORS.items():
        legend_entry(type_name, solid(color), Inches(0.7))
    legend_entry("Pianificato", outlined, Inches(0.85))
    legend_entry("Effettivo", solid(COLOR_GANTT_OTHER), Inches(0.75))
    legend_entry("In corso (fino a oggi)", lambda s: (_fading(s, COLOR_GANTT_OTHER), s.line.fill.background()), Inches(1.5))
    if any(i.is_activity for i in items):
        legend_entry(
            "Attività di servizio",
            lambda s: (_striped(s, COLOR_GANTT_SERVICE_PLANNED), s.line.fill.background()),
            Inches(1.35),
        )
    line(x, legend_top + Inches(0.15), x + Inches(0.25), legend_top + Inches(0.15), COLOR_DANGER, width=1.5)
    cell(x + Inches(0.3), legend_top, Inches(1.3), Inches(0.3), f"Oggi {_date(today)}", size=9, color=COLOR_MUTED)


def _tint(color: RGBColor, amount: float) -> RGBColor:
    """Il colore schiarito verso il bianco: amount = quota del colore
    originale (come color-mix(... amount%, white) nel frontend)."""
    return RGBColor(*(round(c * amount + 255 * (1 - amount)) for c in color))


def _striped(shape, color: RGBColor) -> None:
    """Righe oblique: attivita' di servizio, come nel Gantt del Backlog."""
    shape.fill.patterned()
    shape.fill.pattern = MSO_PATTERN.WIDE_UPWARD_DIAGONAL
    shape.fill.fore_color.rgb = color
    shape.fill.back_color.rgb = _tint(color, 0.35)


def _fading(shape, color: RGBColor) -> None:
    """Pieno a sinistra, sfumato a destra: item in corso, senza ancora una
    fine effettiva."""
    shape.fill.gradient()
    shape.fill.gradient_angle = 0
    stops = shape.fill.gradient_stops
    stops[0].position = 0.7
    stops[0].color.rgb = color
    stops[1].position = 1.0
    stops[1].color.rgb = _tint(color, 0.1)


def _date_serial(day: dt.date) -> int:
    """Data come numero seriale di Excel: l'asse X dei grafici nel tempo."""
    return (day - EXCEL_EPOCH).days


def _date_x_axis(chart, days: list[dt.date]) -> None:
    """Asse X di date per un grafico a dispersione, dalla prima all'ultima
    delle date indicate. In un grafico a dispersione "category_axis" e'
    l'asse X dei valori: le date sono numeri seriali di Excel, mostrati come
    date."""
    x_axis = chart.category_axis
    x_axis.minimum_scale = _date_serial(min(days))
    x_axis.maximum_scale = _date_serial(max(days))
    x_axis.tick_labels.number_format = "dd/mm/yyyy"
    x_axis.tick_labels.number_format_is_linked = False
    x_axis.has_major_gridlines = True
    # Senza questo elemento PowerPoint legge i seriali col sistema di date
    # 1904 e le date escono spostate di 4 anni.
    chart_space = chart._chartSpace
    date1904 = chart_space.makeelement(f"{{{chart_space.nsmap['c']}}}date1904", {"val": "0"})
    chart_space.insert(0, date1904)


def _hours_slide(deck: _Deck, project: models.Project) -> None:
    slide, box = deck.add_slide("Ore usate nel tempo")
    start, freeze = project.start_date, project.code_freeze_date
    if start is None or freeze is None or freeze <= start:
        _text(
            slide,
            box,
            "Servono la data di inizio increment e la data di code freeze (scheda increment) per tracciare le ore nel tempo.",
            color=COLOR_MUTED,
        )
        return

    serial = _date_serial
    chart_data = XyChartData()
    days: list[dt.date] = [start, freeze]
    for name, attribute in (("Ore usate (h)", "actual_hours"), ("Ore Jira (h)", "logged_hours")):
        points = [(s.snapshot_date, getattr(s, attribute)) for s in project.snapshots if getattr(s, attribute) is not None]
        if not points:
            continue
        series = chart_data.add_series(name)
        for day, hours in points:
            series.add_data_point(serial(day), hours)
            days.append(day)
    if len(days) == 2:
        _text(slide, box, "Nessuno snapshot con ore registrate (tab Andamento).", color=COLOR_MUTED)
        return

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER_LINES, box.left, box.top, box.width, box.height, chart_data
    ).chart
    _date_x_axis(chart, days)
    _style_chart(chart, value_axis_title="Ore")
    for series in chart.plots[0].series:
        series.smooth = False
        series.format.line.width = Pt(2.25)


def _completion_slide(deck: _Deck, project: models.Project) -> None:
    slide, box = deck.add_slide("Andamento % completamento nel tempo")
    snapshots = list(project.snapshots)
    if not snapshots:
        _text(
            slide,
            box,
            'Nessuno snapshot registrato. Aggiungine uno dalla tab "Andamento" per iniziare a tracciare lo storico.',
            color=COLOR_MUTED,
        )
        return
    # Come nella Dashboard: % completamento degli snapshot su un asse di date
    # vere e, se inizio e code freeze sono noti, l'obiettivo - la data del
    # code freeze e la linea ideale dallo 0% all'inizio al 100% al code freeze.
    start, freeze = project.start_date, project.code_freeze_date
    has_plan = start is not None and freeze is not None and freeze > start
    points = [
        (s.snapshot_date, round((s.pbi_done or 0) / s.pbi_total * 100)) for s in snapshots if s.pbi_total
    ]
    # Ultimo rilevamento dell'Andamento: il completamento reale e, sulla
    # linea ideale, dove dovremmo essere quel giorno.
    last_day, last_actual = points[-1] if points else (None, None)
    last_ideal = None
    if has_plan and last_day is not None:
        last_ideal = round(min(max((last_day - start).days / (freeze - start).days, 0), 1) * 100)

    chart_data = XyChartData()
    days = [s.snapshot_date for s in snapshots]
    completion = chart_data.add_series("Completamento")
    for day, percent in points:
        completion.add_data_point(_date_serial(day), percent)
    if has_plan:
        days += [start, freeze]
        ideal = chart_data.add_series("Ideale")
        ideal.add_data_point(_date_serial(start), 0)
        # Punto in piu' sulla linea ideale (resta una retta), per poterci
        # scrivere il valore nel giorno dell'ultimo rilevamento.
        if last_ideal is not None and start < last_day < freeze:
            ideal.add_data_point(_date_serial(last_day), last_ideal)
        ideal.add_data_point(_date_serial(freeze), 100)
        deadline = chart_data.add_series(f"Code freeze {_date(freeze)}")
        deadline.add_data_point(_date_serial(freeze), 0)
        deadline.add_data_point(_date_serial(freeze), 100)

    summary_height = 0
    if last_day is not None:
        summary = f"Ultimo rilevamento ({_date(last_day)}): completamento reale {last_actual}%"
        if last_ideal is not None:
            gap = last_actual - last_ideal
            summary += f", ideale {last_ideal}% - " + (
                f"{abs(gap)} punti {'sotto' if gap < 0 else 'sopra'} l'ideale" if gap else "in linea con l'ideale"
            )
        summary_height = Inches(0.45)
        _text(slide, _Box(box.left, box.top, box.width, summary_height), summary + ".", size=14)

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER_LINES,
        box.left,
        box.top + summary_height,
        box.width,
        box.height - summary_height,
        chart_data,
    ).chart
    _date_x_axis(chart, days)
    _style_chart(chart)
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 100
    chart.value_axis.tick_labels.number_format = '0"%"'
    chart.value_axis.tick_labels.number_format_is_linked = False
    plot_series = list(chart.plots[0].series)
    for series in plot_series:
        series.smooth = False
        series.format.line.width = Pt(2.25)
    # Linea ideale e code freeze: tratteggiate e senza punti, fanno da
    # riferimento al dato reale.
    for series, color in zip(plot_series[1:], (COLOR_PLANNED, COLOR_DANGER)):
        series.format.line.color.rgb = color
        series.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        series.marker.style = XL_MARKER_STYLE.NONE

    # I due valori dell'ultimo rilevamento scritti sul grafico: quello piu'
    # alto sopra il suo punto, l'altro sotto, cosi' non si sovrappongono.
    def label(point, text: str, color: RGBColor, position) -> None:
        data_label = point.data_label
        data_label.position = position
        _run(data_label.text_frame.paragraphs[0], text, size=12, bold=True, color=color)

    if last_day is not None:
        actual_above = last_ideal is None or last_actual >= last_ideal
        label(
            plot_series[0].points[len(points) - 1],
            f"Reale {last_actual}%",
            COLOR_TEXT,
            XL_LABEL_POSITION.ABOVE if actual_above else XL_LABEL_POSITION.BELOW,
        )
        if last_ideal is not None and start < last_day < freeze:
            # Di lato e non sopra/sotto: la linea ideale sale verso destra,
            # un'etichetta centrata sul punto ci finirebbe sopra.
            label(
                plot_series[1].points[1],
                f"Ideale {last_ideal}%",
                COLOR_MUTED,
                XL_LABEL_POSITION.RIGHT if actual_above else XL_LABEL_POSITION.LEFT,
            )


def _plan_vs_actual_slides(
    deck: _Deck,
    project: models.Project,
    title: str,
    empty_message: str,
    planned_label: str,
    actual_label: str,
    unit: str,
    get_values,
) -> None:
    """Per ogni PBI in scope (nell'ordine del Backlog) previsione, dato
    effettivo e scostamento effettivo - previsto: rosso se oltre la
    previsione, verde se entro. Esclusi i PBI a cui manca uno dei due valori."""
    points: list[tuple[str, float, float]] = []
    for item in sorted((i for i in project.backlog_items if i.in_scope), key=lambda i: i.priority_order):
        planned, actual = get_values(item)
        if planned is not None and actual is not None:
            points.append((item.jira_key, planned, actual))

    if not points:
        slide, box = deck.add_slide(title)
        _text(slide, box, empty_message, color=COLOR_MUTED)
        return

    over = sum(1 for _, planned, actual in points if actual > planned)
    total_planned = sum(planned for _, planned, _ in points)
    total_actual = sum(actual for _, _, actual in points)
    summary = (
        f"{len(points)} PBI: {over} oltre la previsione, {len(points) - over} entro. "
        f"Totale {_num(total_actual)} {unit} effettivi su {_num(total_planned)} {unit} previsti"
    )
    if total_planned > 0:
        delta_pct = round((total_actual - total_planned) / total_planned * 100)
        summary += f" ({'+' if delta_pct > 0 else ''}{delta_pct}%)"
    summary += "."

    pages = [
        points[i : i + PLAN_VS_ACTUAL_POINTS_PER_SLIDE] for i in range(0, len(points), PLAN_VS_ACTUAL_POINTS_PER_SLIDE)
    ]
    for index, page in enumerate(pages):
        slide, box = deck.add_slide(title if len(pages) == 1 else f"{title} - {index + 1}/{len(pages)}")
        summary_height = Inches(0.45)
        _text(slide, _Box(box.left, box.top, box.width, summary_height), summary, size=14)

        chart_data = CategoryChartData()
        chart_data.categories = [key for key, _, _ in page]
        chart_data.add_series(planned_label, [planned for _, planned, _ in page])
        chart_data.add_series(actual_label, [actual for _, _, actual in page])
        chart_data.add_series("Scostamento", [actual - planned for _, planned, actual in page])
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            box.left,
            box.top + summary_height,
            box.width,
            box.height - summary_height,
            chart_data,
        ).chart
        _style_chart(chart, value_axis_title=unit)
        plot = chart.plots[0]
        plot.gap_width = 60
        plot.overlap = -5
        planned_series, actual_series, delta_series = plot.series
        planned_series.format.fill.solid()
        planned_series.format.fill.fore_color.rgb = COLOR_PLANNED
        actual_series.format.fill.solid()
        actual_series.format.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        delta_series.invert_if_negative = False
        delta_series.format.fill.solid()
        delta_series.format.fill.fore_color.rgb = COLOR_DANGER
        for n, (_, planned, actual) in enumerate(page):
            fill = delta_series.points[n].format.fill
            fill.solid()
            fill.fore_color.rgb = COLOR_DANGER if actual > planned else COLOR_SUCCESS if actual < planned else COLOR_MUTED
            # Il colore del singolo punto vale anche sotto lo zero solo se
            # "inverti se negativo" e' spento anche sul punto (altrimenti la
            # barra negativa esce bianca).
            point = delta_series._element.get_or_add_dPt_for_point(n)
            invert = point.makeelement(f"{{{point.nsmap['c']}}}invertIfNegative", {"val": "0"})
            point.insert(1, invert)

        axis = chart.category_axis
        # Etichette in basso anche con barre negative (scostamenti entro la
        # previsione), e inclinate quando i PBI sono molti.
        axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        axis.tick_labels.font.size = Pt(9)
        if len(page) > 12:
            axis._element.xpath("c:txPr/a:bodyPr")[0].set("rot", "-2700000")


def generate_dashboard_presentation(project: models.Project, closed_days: int = 7) -> tuple[bytes, str]:
    """Presentazione della Dashboard dell'increment. closed_days: il periodo
    della card "Issue chiuse" selezionato nella pagina. Restituisce il
    contenuto del .pptx e il nome file (senza estensione)."""
    today = dt.date.today()
    deck = _Deck(footer=f"{project.code} - {project.name}   ·   {_date(today)}")
    deck.add_cover(
        f"{project.code} - {project.name}",
        ["Dashboard increment", f"Stato: {project.status}", _date(today)],
    )
    _scope_slides(deck, project)
    _kpi_slide(deck, project)
    _recently_closed_slides(deck, project, closed_days)
    _phases_slide(deck, project)
    _timeline_slide(deck, project)
    _backlog_gantt_slides(deck, project)
    _hours_slide(deck, project)
    _completion_slide(deck, project)
    _plan_vs_actual_slides(
        deck,
        project,
        title="Sizing vs Durata",
        empty_message=(
            "Nessun PBI con sia il Sizing (gg) sia la Durata (gg): servono il Sizing compilato nel Backlog e le "
            "date effettive di inizio/fine dal sync Jira."
        ),
        planned_label="Sizing",
        actual_label="Durata",
        unit="gg",
        get_values=lambda i: (i.planned_duration_days, _working_days_between(i.actual_start, i.actual_finish)),
    )
    _plan_vs_actual_slides(
        deck,
        project,
        title="Ore stimate vs Ore Jira",
        empty_message="Nessun PBI con sia le Ore stimate sia le Ore Jira.",
        planned_label="Ore stimate",
        actual_label="Ore Jira",
        unit="h",
        get_values=lambda i: (i.dev_effort_hours, i.logged_hours),
    )
    # Ora nel nome: due presentazioni generate lo stesso giorno non si
    # sovrascrivono (in una cartella OneDrive, salvare una versione sopra
    # l'altra puo' farle fondere da PowerPoint, con le forme duplicate).
    return deck.to_bytes(), f"{project.code} - Dashboard {dt.datetime.now():%Y-%m-%d %H%M}"
