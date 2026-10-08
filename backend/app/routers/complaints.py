import datetime as dt
import re
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app import models, schemas
from app.database import get_db
from app.routers.settings import SETTINGS_ROW_ID
from app.services.jira_client import JiraClientError, fetch_complaints

router = APIRouter(prefix="/api/complaints", tags=["complaints"])

# Architettura e Sito Cliente non hanno un campo Jira dedicato: sono label
# dell'issue. L'architettura e' una di queste quattro (stesso elenco di
# schemas.ComplaintUpdate.architecture); se la label manca o e' sbagliata si
# imposta a mano dalla pagina (vedi Complaint.architecture_manual).
ARCHITECTURES = ["Legacy", "NA5", "NA6", "NA7"]

# Il Sito Cliente e' una label "libera" (es. "Careggi", "Policlinico_Milano"),
# mescolata alle label di lavorazione: si considerano sito le label che
# iniziano con una maiuscola, non contengono cifre (via release come
# "PTBSYS-02-017", "CAPA30"...) e non sono in questo elenco di label note
# che non sono siti. Dove il risultato e' sbagliato il sito si corregge a
# mano dalla pagina (vedi Complaint.customer_site_manual).
NON_SITE_LABELS = {
    "Analysis",
    "Analytix",
    "Board_Redesign",
    "DAS",
    "Development",
    "HP",
    "IPTB",
    "PDP",
    "PTB-Maintenance",
    "PTB_Master_HW&SW",
    "Postpone-Checkout",
    "Postponed",
    "TSB",
    "Transportation",
}

JIRA_BROWSE_URL = "https://inpeco.atlassian.net/browse/"

# Colonne dell'export Excel: le stesse della tabella della pagina, nello
# stesso ordine (intestazione, larghezza della colonna).
EXPORT_COLUMNS = [
    ("ID Salesforce", 15),
    ("Summary", 80),
    ("Stato Salesforce", 18),
    ("ID Jira", 15),
    ("Stato Jira", 15),
    ("Fix Version", 28),
    ("Severity", 12),
    ("Creato su Jira", 16),
    ("Architettura", 15),
    ("Sito Cliente", 28),
]

# Il numero del case Salesforce e' anche in testa al summary ("[76665]...").
_SUMMARY_CASE_RE = re.compile(r"^\s*\[(\d+)\]")


def _architecture(labels: list[str]) -> str | None:
    return ", ".join(a for a in ARCHITECTURES if a in labels) or None


def _customer_site(labels: list[str]) -> str | None:
    sites = [
        label.replace("_", " ")
        for label in labels
        if label not in ARCHITECTURES
        and label not in NON_SITE_LABELS
        and label[:1].isupper()
        and not any(ch.isdigit() for ch in label)
    ]
    return ", ".join(sites) or None


def _case_number(source_note: str | None, summary: str) -> str | None:
    if source_note and source_note.strip():
        return source_note.strip()
    match = _SUMMARY_CASE_RE.match(summary or "")
    return match.group(1) if match else None


def _settings_row(db: Session) -> models.AppSettings | None:
    return db.get(models.AppSettings, SETTINGS_ROW_ID)


def _base_jql(row: models.AppSettings | None) -> str:
    return (row.complaints_base_jql if row else None) or models.DEFAULT_COMPLAINTS_JQL


@router.get("", response_model=list[schemas.Complaint])
def list_complaints(db: Session = Depends(get_db)):
    return (
        db.query(models.Complaint)
        .order_by(models.Complaint.jira_created.desc(), models.Complaint.jira_key.desc())
        .all()
    )


# Dichiarate prima di "/{complaint_id}": altrimenti "settings", "export" e
# "sync" verrebbero lette come id.
@router.get("/settings", response_model=schemas.ComplaintsSettings)
def get_complaints_settings(db: Session = Depends(get_db)):
    return schemas.ComplaintsSettings(base_jql=_base_jql(_settings_row(db)))


@router.get("/export")
def export_complaints(db: Session = Depends(get_db)):
    """Tutti i complaint in un foglio Excel (indipendentemente dai filtri
    attivi nella pagina), dal piu' nuovo, con le intestazioni filtrabili."""
    complaints = (
        db.query(models.Complaint)
        .order_by(models.Complaint.jira_created.desc(), models.Complaint.jira_key.desc())
        .all()
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Complaints"
    sheet.append([header for header, _ in EXPORT_COLUMNS])
    for c in complaints:
        sheet.append(
            [
                c.salesforce_case_number,
                c.summary,
                c.salesforce_status,
                c.jira_key,
                c.jira_status,
                c.fix_versions,
                c.severity,
                c.jira_created,
                c.architecture,
                c.customer_site,
            ]
        )

    for index, (_, width) in enumerate(EXPORT_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        header = sheet.cell(row=1, column=index)
        header.font = Font(bold=True, color="FFFFFF")
        header.fill = PatternFill("solid", fgColor="2F6FED")
        header.alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == 2)
        key = row[3]
        key.hyperlink = f"{JIRA_BROWSE_URL}{key.value}"
        key.font = Font(color="0563C1", underline="single")
        row[7].number_format = "DD/MM/YYYY"
    # Filtro automatico su tutte le colonne, con l'intestazione bloccata in
    # alto scorrendo le righe.
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"

    buffer = BytesIO()
    workbook.save(buffer)
    filename = f"Complaints {dt.datetime.now():%Y-%m-%d %H%M}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.put("/settings", response_model=schemas.ComplaintsSettings)
def update_complaints_settings(payload: schemas.ComplaintsSettings, db: Session = Depends(get_db)):
    row = _settings_row(db)
    if row is None:
        row = models.AppSettings(id=SETTINGS_ROW_ID)
        db.add(row)
    # Vuota = si torna alla JQL di default.
    row.complaints_base_jql = payload.base_jql.strip() or None
    db.commit()
    return schemas.ComplaintsSettings(base_jql=_base_jql(row))


@router.post("/sync", response_model=schemas.SyncResult)
def sync_complaints_from_jira(db: Session = Depends(get_db)):
    row = _settings_row(db)
    try:
        issues = fetch_complaints(
            (row.jira_base_url if row else None) or "",
            (row.jira_email if row else None) or "",
            (row.jira_api_token if row else None) or "",
            _base_jql(row),
        )
    except JiraClientError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    # Una JQL sbagliata (o troppo stretta) che non trova nulla svuoterebbe la
    # tabella, perdendo gli stati Salesforce e i siti inseriti a mano: in quel
    # caso non si tocca niente.
    if not issues:
        raise HTTPException(
            status_code=400,
            detail="La JQL non ha restituito nessuna issue: sincronizzazione annullata, i complaint esistenti non sono stati modificati.",
        )

    existing = {c.jira_key: c for c in db.query(models.Complaint)}
    matched_keys = {issue.key for issue in issues}
    now = dt.datetime.utcnow()
    created = 0
    updated = 0

    for issue in issues:
        complaint = existing.get(issue.key)
        if complaint is None:
            # salesforce_status resta al default ("Aperto"): non e' in Jira.
            complaint = models.Complaint(jira_key=issue.key)
            db.add(complaint)
            created += 1
        else:
            updated += 1
        complaint.summary = issue.summary
        complaint.jira_status = issue.status
        complaint.fix_versions = ", ".join(issue.fix_versions) or None
        complaint.severity = issue.severity
        complaint.jira_created = issue.created
        complaint.jira_resolved = issue.resolved
        complaint.labels = ";".join(issue.labels)
        if not complaint.architecture_manual:
            complaint.architecture = _architecture(issue.labels)
        complaint.salesforce_case_number = _case_number(issue.source_note, issue.summary)
        complaint.salesforce_case_id = issue.salesforce_case_id
        if not complaint.customer_site_manual:
            complaint.customer_site = _customer_site(issue.labels)
        complaint.last_synced_at = now

    # Come per il Backlog: cio' che non e' piu' nel risultato della JQL (es.
    # Source Type cambiato su Jira) viene rimosso.
    removed = 0
    for jira_key, complaint in existing.items():
        if jira_key not in matched_keys:
            db.delete(complaint)
            removed += 1

    db.commit()
    return schemas.SyncResult(created=created, updated=updated, removed=removed, total_matched=len(issues))


@router.put("/{complaint_id}", response_model=schemas.Complaint)
def update_complaint(complaint_id: int, payload: schemas.ComplaintUpdate, db: Session = Depends(get_db)):
    complaint = db.get(models.Complaint, complaint_id)
    if complaint is None:
        raise HTTPException(status_code=404, detail="Complaint non trovato")
    data = payload.model_dump(exclude_unset=True)
    if data.get("salesforce_status"):
        complaint.salesforce_status = data["salesforce_status"]
    labels = (complaint.labels or "").split(";")
    if "architecture" in data:
        architecture = data["architecture"] or None
        complaint.architecture_manual = architecture is not None
        complaint.architecture = architecture or _architecture(labels)
    if "customer_site" in data:
        site = (data["customer_site"] or "").strip()
        complaint.customer_site_manual = bool(site)
        complaint.customer_site = site or _customer_site(labels)
    db.commit()
    db.refresh(complaint)
    return complaint
